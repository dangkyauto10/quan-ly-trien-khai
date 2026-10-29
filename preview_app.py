# QUAN LY DU AN — PREVIEW 2026-10-22
# Staged only: never deploy to production without user approval.
import io
import json
import base64
from pathlib import Path
import uuid
import hmac
import hashlib
import time
import threading
import unicodedata
from preview_site_search import search_assigned_site
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from urllib.parse import quote, urlparse

import gspread
import streamlit as st
import pandas as pd
from google.oauth2.service_account import Credentials
from google.oauth2.credentials import Credentials as OAuthCredentials
from google.auth.transport.requests import Request
from preview_drive_picker import pick_folder
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload
from PIL import Image
from streamlit_geolocation import streamlit_geolocation
# The cookie component runs once on each page execution. Its browser-side
# cookie cache is populated asynchronously after the component has mounted.
from streamlit_cookies_controller import CookieController
from preview_remember import (
    COOKIE_NAME, VALID_SECONDS, issue_remember_token, verify_remember_token
)

from preview_registration import role_for_specialty, team_id_from_registration, matching_teams
from preview_approval_ui import render_pending_approvals
from preview_assignment import (
    assignment_rows, assignment_batch_requests, allowed_modes,
    eligible_site_tasks, project_site_catalog, site_assignment_plan, compact_site_assignment_requests
)
from preview_bulk_allocation import (
    inventory_by_sku, allocation_waiting_summary,
    build_plan, verify_plan, atomic_requests
)
from preview_bulk_excel import parse_exception_excel
from preview_allocation_adjust import prepare_quantity_corrections, allocation_batch_code
from preview_allocation_groups import (
    match_contract_locations, read_contract_locations_excel, preview_basket
)
from preview_leadership import build_progress, display_amount
from preview_leader_access import (
    LEADER_SESSION_USERNAME, check_leader_password,
    make_leader_session, verify_leader_session
)
from preview_team_access import resolve_login_username, friendly_account_name, project_login_choices
from preview_team_access import normalized_name
from preview_credential_accounts import (
    team_accounts, visible_credential_state, next_account_id,
    new_password, create_hash, validate_custom_password
)
from preview_signup_credentials import registration_requests, pending_row
from preview_team_detach import (
    lost_registrations, revoked_team_ids, plan_detach,
    detach_requests, ACTIVE as LINK_ACTIVE
)
from preview_team_reassignment import (
    team_change_plan, change_signature, team_change_requests
)
from preview_site_rights import (
    allocation_owners, active_registration, site_ids, RIGHT_FIELDS,
    owner_index, supplemental_site_plan, supplemental_site_signature,
    supplemental_site_requests
)
from preview_assignment import ROLE_TABS, TASK_SCHEMA
from preview_remember import account_fingerprint
from preview_handover import installer_handover, site_may_complete_installation
from preview_report_commit import (atomic_report_requests, pending_quantity, numeric,
                                   allocated_report_item)
from preview_progress_rules import site_progress
from preview_hub import hub_plan, received_by_site, signed_receipt_lines, qty_label, verify_vc_assignment
from preview_user_guide import GUIDE_URL, render_guide
from preview_project_cleanup import project_cleanup_admin
from io import BytesIO
from openpyxl import Workbook as ExcelWorkbook
from preview_import_locations import (read_location_file, candidate_names,
    column_choice, is_header, normalize_name, blank_excel_template)


def display_centered_table(data, *args, **kwargs):
    """Display read-only tables with centered bold headers and optional compact columns."""
    center_columns = tuple(kwargs.pop("center_columns", ()))
    compact = bool(kwargs.pop("compact", False))
    frame = data if isinstance(data, pd.DataFrame) else pd.DataFrame(data)
    # st.dataframe draws headers on a canvas and ignores Pandas header CSS.
    # These read-only summaries use escaped HTML so header styling is reliable.
    if args or any(key not in {"hide_index", "use_container_width", "height", "width"}
                   for key in kwargs):
        return st.dataframe(frame, *args, **kwargs)
    table_id = "tdv-table-" + uuid.uuid4().hex
    height = kwargs.get("height")
    hide_index = kwargs.get("hide_index", False)
    scrolling = ("max-height:" + str(max(80, int(height))) + "px;"
                 if isinstance(height, (int, float)) else "")
    markup = frame.to_html(index=not hide_index, escape=True, border=0, na_rep="")
    table_width = ("width:fit-content;max-width:100%;table-layout:auto;"
                   if compact else "width:100%;")
    body_center_rules = "".join(
        "#" + table_id + " tbody td:nth-child(" +
        str(index + 1 + (0 if hide_index else 1)) +
        "){text-align:center!important;}"
        for index, column in enumerate(frame.columns)
        if column in center_columns
    )
    compact_column_rules = (
        "#" + table_id + " tbody td:nth-child(2),#" + table_id +
        " thead th:nth-child(2){max-width:min(58vw,620px);"
        "white-space:normal;overflow-wrap:anywhere;}"
        if compact else ""
    )
    return st.markdown(
        "<style>"
        "#" + table_id + "{overflow:auto;" + scrolling + "}"
        "#" + table_id + " table{" + table_width +
        "border-collapse:collapse;font-size:0.9rem;}"
        "#" + table_id + " th,#" + table_id + " td{"
        "padding:0.45rem 0.6rem;border-bottom:1px solid #dfe3e8;}"
        "#" + table_id + " thead th{"
        "text-align:center!important;font-weight:700!important;"
        "background:var(--secondary-background-color,#f0f2f6);"
        "position:sticky;top:0;z-index:1;}"
        "#" + table_id + " tbody td,#" + table_id + " tbody th{"
        "text-align:left;font-weight:500!important;}"
        + body_center_rules + compact_column_rules +
        "</style><div id=\"" + table_id + "\">" + markup + "</div>",
        unsafe_allow_html=True,
    )

_OVERVIEW_ROLE_MODES = {
    "VC": ("VAN_CHUYEN",),
    "KTV": ("LAP_DAT",),
    "VC+KTV": ("VAN_CHUYEN", "LAP_DAT"),
}

def approved_team_site_overview(team, registrations, locations, tasksets):
    """Display already-approved site rights even before SKU allocation.

    This is read-only. Actual VC/KTV reports still require allocated tasks.
    """
    project = str(team.get("Mã dự án") or "").strip()
    team_id = str(team.get("Mã đội") or "").strip()
    role = team.get("Loại nhiệm vụ")
    if (not project or not team_id or role not in _OVERVIEW_ROLE_MODES or
            team.get("Trạng thái") != "Hoạt động"):
        raise ValueError("Đội chưa hoạt động hoặc thiếu mã dự án/chuyên môn.")
    registration = active_registration(registrations, project, team_id)
    project_sites = {r.get("Mã điểm"): r for r in locations
                     if r.get("Mã dự án") == project and r.get("Mã điểm")}
    overview = []
    for tab in _OVERVIEW_ROLE_MODES[role]:
        field, _ = RIGHT_FIELDS[tab]
        _, assigned_field, _ = TASK_SCHEMA[tab]
        for site_id in site_ids(registration.get(field)):
            site = project_sites.get(site_id)
            if site is None:
                raise ValueError("Địa điểm đã duyệt không còn thuộc dự án: " + site_id)
            tasks = [r for r in tasksets.get(tab, [])
                     if r.get("Mã dự án") == project and
                     r.get("Mã điểm") == site_id]
            linked = sum(r.get(assigned_field) == team_id for r in tasks)
            if site.get("Trạng thái") != "Hoạt động":
                status = "Địa điểm tạm dừng – cần Admin kiểm tra"
            elif not tasks:
                status = "Đã duyệt điểm – Chưa phân bổ SKU"
            elif linked == len(tasks):
                status = "Đã duyệt điểm – Đã có nhiệm vụ SKU"
            else:
                status = "Đã duyệt điểm – Cần đối soát nhiệm vụ SKU"
            overview.append({
                "site_id": site_id,
                "site_name": site.get("Tên xã/phường/đơn vị") or site_id,
                "role": "Vận chuyển" if tab == "VAN_CHUYEN" else "Lắp đặt",
                "sku_tasks": len(tasks),
                "linked_tasks": linked,
                "status": status,
            })
    return overview



APP_ICON = Path(__file__).resolve().parent / "assets" / "tdv_icon_144.png"
# Use the approved TDV square logo for the tab favicon and app masthead.
st.set_page_config(
    page_title="DỰ ÁN",
    page_icon=str(APP_ICON),
    layout="wide"
)
APP_ICON_DATA = base64.b64encode(APP_ICON.read_bytes()).decode("ascii")
# Set the actual host page metadata (Streamlit renders the Python app
# inside a same-origin iframe; window.parent alone only changes the inner page).
import streamlit.components.v1 as components
components.html("""
<script>
(() => {
  let doc;
  try { doc = window.top.document; } catch (_) { return; }
  const iconSource = doc.head.querySelector('link[rel="alternate icon"]')
    || doc.head.querySelector('link[rel~="icon"]');
  const iconUrl = iconSource ? iconSource.href : "";
  const setLink = (rel, href) => {
    let link = doc.head.querySelector('link[rel="' + rel + '"]');
    if (!link) { link = doc.createElement("link"); link.rel = rel; doc.head.appendChild(link); }
    link.href = href;
    if (rel === "manifest") {
      link.type = "application/manifest+json";
      link.crossOrigin = "anonymous";
    } else {
      link.type = "image/png";
      link.sizes = "144x144";
    }
  };
  if (iconUrl) {
    setLink("apple-touch-icon", iconUrl);
    setLink("icon", iconUrl);
  }
  // Override Streamlit Cloud's default manifest name, which is "Streamlit".
  setLink("manifest", "/~/+/app/static/manifest.json?v=20261027h");
  const setMeta = (selector, attribute, value) => {
    let meta = doc.head.querySelector(selector);
    if (!meta) { meta = doc.createElement("meta"); meta.setAttribute(attribute, selector.includes("property=") ? selector.match(/property="([^"]+)"/)[1] : selector.match(/name="([^"]+)"/)[1]); doc.head.appendChild(meta); }
    meta.content = value;
  };
  setMeta('meta[name="apple-mobile-web-app-title"]', "name", "DỰ ÁN");
  setMeta('meta[name="application-name"]', "name", "DỰ ÁN");
  setMeta('meta[property="og:title"]', "property", "DỰ ÁN");
  doc.title = "DỰ ÁN";
})();
</script>
""", width=0, height=0)


cookie_controller = CookieController(key="tdv_cookie_controller_v1")


def remember_signing_key():
    raw = st.secrets.get("remember_login_secret", "")
    return raw if isinstance(raw, str) and len(raw) >= 32 else None


def sync_remembered_login():
    """Handle a pending browser-cookie change before trying auto login."""
    action = st.session_state.pop("_remember_action", None)
    if action == "set":
        token = st.session_state.pop("_remember_token", "")
        if token:
            cookie_controller.set(
                COOKIE_NAME, token,
                max_age=VALID_SECONDS, expires=datetime.now() + timedelta(days=30),
                secure=True, same_site="strict")
    elif action == "remove":
        st.session_state.pop("_remember_token", None)
        if cookie_controller.get(COOKIE_NAME) is not None:
            cookie_controller.remove(COOKIE_NAME)

    if not st.session_state.get("auth") and action != "remove":
        key = remember_signing_key()
        cookie = cookie_controller.get(COOKIE_NAME)
        if key and cookie:
            accounts=effective_accounts()
            account = verify_remember_token(cookie, accounts, key)
            if account:
                account["_fp"]=account_fingerprint(accounts[account["username"]])
                st.session_state["auth"] = account



st.markdown("""
<style>
.block-container {
  /* Wide PC workspace for admin tables and project dashboards. */
  max-width: 100% !important;
  width: 100% !important;
  padding: 2.1rem clamp(1rem, 2vw, 2.5rem) 3rem !important;
}
/* Preserve the already-approved compact tablet/phone experience. */
@media (max-width: 900px) {
  .block-container {
    max-width: 720px !important;
    margin-left: auto !important;
    margin-right: auto !important;
  }
}
.tdv-hero {
  box-sizing: border-box;
  width: 100%;
  text-align: center;
  /* Reserve vertical room for Vietnamese accents on all devices.
     Streamlit may otherwise clip the top of a short HTML line. */
  padding: 1.4rem 0.3rem 1rem;
  margin: 0 auto 0.55rem;
  height: auto !important;
  min-height: 9rem;
  overflow: visible !important;
}
.tdv-hero-title {
  font-family: Arial, Helvetica, sans-serif !important;
  display: flex;
  align-items: center;
  justify-content: center;
  gap: 0.5rem;
  color: #173A73;
  font-weight: 800;
  font-size: clamp(1.55rem, 4.5vw, 2rem);
  line-height: 1.75 !important;
  letter-spacing: 0;
  white-space: normal;
  flex-wrap: wrap;
  min-height: 4rem;
  padding: 0.25rem 0;
  overflow: visible !important;
}
.tdv-hero-name {
  font-family: Arial, Helvetica, sans-serif !important;
  display: inline-block;
  line-height: 1.8 !important;
  letter-spacing: 0;
  padding: 0.12rem 0;
  overflow: visible !important;
}
.tdv-hero-icon {
  display: block;
  width: 3.15rem;
  height: 3.15rem;
  object-fit: contain;
  flex-shrink: 0;
  border-radius: 0.68rem;
}
.tdv-hero-subtitle {
  margin-top: 0.18rem;
  color: #64748b;
  font-size: 0.84rem;
  line-height: 1.6;
}
.tdv-preview-tag {
  display: inline-block;
  margin-top: 0.5rem;
  padding: 0.22rem 0.7rem;
  border: 1px solid #dbe5f2;
  background: #f5f8fc;
  border-radius: 999px;
  color: #526581;
  font-size: 0.66rem;
  line-height: 1.5;
  letter-spacing: 0.02em;
}
@media (max-width: 600px) {
  /* Reserve space for Streamlit Cloud's fixed top bar; keep the brand below it. */
  .block-container {
    padding: 4.5rem 0.65rem 1.25rem !important;
    font-family: system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif !important;
  }
  .stApp button, .stApp input, .stApp textarea, .stApp label,
  .stApp select, .stApp [data-testid="stMarkdownContainer"] {
    font-family: system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif !important;
  }
  .tdv-hero {
    position: sticky;
    top: 4rem;
    z-index: 100;
    background: rgba(255, 255, 255, 0.98);
    padding: 0.22rem 0.25rem 0.28rem;
    margin-bottom: 0.25rem;
    min-height: 0;
    border-bottom: 1px solid #edf2f7;
  }
  .tdv-hero-title {
    font-size: clamp(1.08rem, 4.8vw, 1.32rem);
    font-weight: 700;
    gap: 0.28rem;
    line-height: 1.45 !important;
    min-height: 0;
    padding: 0.08rem 0;
  }
  .tdv-hero-name {
    line-height: 1.45 !important;
    padding: 0.04rem 0;
  }
  .tdv-hero-icon {width: 1.8rem; height: 1.8rem;}
  .tdv-hero-subtitle {display: none;}
  .tdv-preview-tag {
    margin-top: 0.08rem;
    padding: 0.06rem 0.36rem;
    font-size: 0.66rem !important;
  }
  .stApp [data-testid="stVerticalBlock"] {gap: 0.65rem;}
  .stApp [data-testid="stHorizontalBlock"] {gap: 0.5rem;}
  div[data-testid="stMetric"] {
    padding: 8px 10px;
    border-radius: 10px;
  }
  div[data-testid="stForm"] {
    padding: 10px;
    border-radius: 12px;
  }
  .stApp div.stButton > button,
  .stApp div.stLinkButton > a {
    min-height: 44px !important;
    font-size: 1rem !important;
  }
}
h2, h3 {color: #173A73;}
div[data-testid="stMetric"] {
  border: 1px solid #DFEAF5;
  background: #F2F7FC;
  border-radius: 14px;
  padding: 12px 14px;
}
div[data-testid="stForm"] {
  border: 1px solid #D9E5F2;
  border-radius: 16px;
  background: #F8FBFE;
  padding: 16px;
}
div.stButton > button, div.stLinkButton > a {
  border-radius: 11px;
  min-height: 43px;
  font-weight: 700;
}
div.stButton > button[kind="primary"] {
  background: #155EEF;
  color: white;
}
div[data-baseweb="select"] > div {border-radius: 10px;}
input {border-radius: 10px !important;}
/* TDV PREVIEW - TANG CO CHU DE DOC */
.tdv-hero-subtitle {
  font-size: 1.04rem !important;
}
.tdv-preview-tag {
  font-size: 0.95rem !important;
}
.stApp [data-testid="stWidgetLabel"] p,
.stApp [data-testid="stMarkdownContainer"] p,
.stApp [data-testid="stMarkdownContainer"] li,
.stApp label {
  font-size: 1.1rem !important;
  line-height: 1.6 !important;
}
.stApp [data-testid="stCaptionContainer"] p {
  font-size: 1rem !important;
}
.stApp input,
.stApp textarea,
.stApp div[data-baseweb="select"] > div,
.stApp div.stButton > button {
  font-size: 1.1rem !important;
}
.stApp div.stButton > button {
  min-height: 48px !important;
}
/* Compact iPhone typography after the app-wide accessibility defaults above. */
@media (max-width: 600px) {
  .tdv-preview-tag {font-size: 0.66rem !important;}
  .stApp [data-testid="stWidgetLabel"] p,
  .stApp [data-testid="stMarkdownContainer"] p,
  .stApp [data-testid="stMarkdownContainer"] li,
  .stApp label {
    font-size: 0.98rem !important;
    line-height: 1.45 !important;
  }
  .stApp [data-testid="stCaptionContainer"] p {font-size: 0.88rem !important;}
  .stApp input,
  .stApp textarea,
  .stApp div[data-baseweb="select"] > div,
  .stApp div.stButton > button {
    font-size: 0.98rem !important;
  }
  .stApp div.stButton > button {min-height: 44px !important;}
}
</style>
""", unsafe_allow_html=True)

st.markdown(f"""
<div class="tdv-hero">
  <div class="tdv-hero-title">
    <img class="tdv-hero-icon" src="data:image/png;base64,{APP_ICON_DATA}" alt="TDV" />
    <span class="tdv-hero-name" lang="vi">QU&#7842;N L&#221; D&#7920; &#193;N</span>
  </div>
  <div class="tdv-hero-subtitle">Hệ thống điều hành hiện trường</div>
  <div class="tdv-preview-tag">BẢN THỬ NGHIỆM · CHƯA VẬN HÀNH</div>
</div>
""", unsafe_allow_html=True)


SHEETS_SCOPE = ["https://www.googleapis.com/auth/spreadsheets",
                "https://www.googleapis.com/auth/drive.file"]
DRIVE_UPLOAD_SCOPE = ["https://www.googleapis.com/auth/drive.file"]
READY = "Đã đồng bộ"

class PreviewSheetReadError(RuntimeError):
    """A read-only Google Sheets request failed; safe error for the UI."""


def _google_api_status(error):
    return getattr(getattr(error, "response", None), "status_code", None)


def _google_read_retryable(error):
    code = _google_api_status(error)
    if code in (429, 500, 502, 503, 504):
        return True
    if code != 403:
        return False
    try:
        payload = error.response.json().get("error", {})
        return any(item.get("reason") in ("rateLimitExceeded", "userRateLimitExceeded")
                   for item in payload.get("errors", []))
    except (AttributeError, ValueError, TypeError):
        return False


def _retry_google_read(operation):
    """Retry only transient READ failures; never replay a write/delete."""
    for attempt in range(4):
        try:
            return operation()
        except gspread.exceptions.APIError as error:
            if attempt == 3 or not _google_read_retryable(error):
                raise
            time.sleep(2 ** attempt)


def now():
    return datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).isoformat(timespec="seconds")

def sheet_datetime_serial():
    """Store a real Sheets date-time (Vietnam local time) for formatted columns."""
    local = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None, microsecond=0)
    return (local - datetime(1899, 12, 30)).total_seconds() / 86400

def fail(message):
    st.error(message)
    st.stop()

@st.cache_resource
def client():
    if "spreadsheet_id" not in st.secrets:
        raise RuntimeError("Thiếu spreadsheet_id trong Streamlit Secrets.")
    # Support the original JSON file pasted as a literal TOML multiline value.
    if "gcp_service_account_json" in st.secrets:
        info = json.loads(st.secrets["gcp_service_account_json"])
    elif "gcp_service_account" in st.secrets:
        info = dict(st.secrets["gcp_service_account"])
    else:
        raise RuntimeError("Thiếu gcp_service_account_json trong Streamlit Secrets.")
    cred = Credentials.from_service_account_info(info, scopes=SHEETS_SCOPE)
    gs = gspread.authorize(cred)
    sh = _retry_google_read(lambda: gs.open_by_key(st.secrets["spreadsheet_id"]))
    return sh, cred


def drive_oauth_credentials():
    """Build the narrowly scoped owner OAuth credentials."""
    oauth = st.secrets.get("google_drive_oauth", {})
    required = ("client_id", "client_secret", "refresh_token")
    missing = [key for key in required if not str(oauth.get(key, "")).strip()]
    if missing:
        raise RuntimeError(
            "Thiếu [google_drive_oauth] client_id, client_secret hoặc "
            "refresh_token trong Streamlit Secrets.")
    cred = OAuthCredentials(
        token=None,
        refresh_token=str(oauth["refresh_token"]).strip(),
        token_uri="https://oauth2.googleapis.com/token",
        client_id=str(oauth["client_id"]).strip(),
        client_secret=str(oauth["client_secret"]).strip(),
        scopes=DRIVE_UPLOAD_SCOPE,
    )
    if not cred.valid:
        cred.refresh(Request())
    return cred


def drive_client():
    """Use only the owner's per-file Drive OAuth grant."""
    cred = drive_oauth_credentials()
    return build("drive", "v3", credentials=cred, cache_discovery=False)


def verify_configured_photo_folder(credentials=None):
    """Verify the configured folder and report the OAuth identity on failure."""
    folder = str(st.secrets.get("photos_folder_id", "")).strip()
    if not folder:
        raise RuntimeError("Thiếu photos_folder_id trong Streamlit Secrets.")
    if credentials is None and st.session_state.get(
            "_tdv_verified_photo_folder_id") == folder:
        return {"id": folder, "name": folder,
                "mimeType": "application/vnd.google-apps.folder"}
    cred = credentials or drive_oauth_credentials()
    drive = build("drive", "v3", credentials=cred, cache_discovery=False)
    oauth_email = ""
    try:
        about = drive.about().get(fields="user(emailAddress)").execute()
        oauth_email = str((about.get("user") or {}).get("emailAddress") or "").strip()
    except Exception:
        # Folder access check remains authoritative if profile metadata is unavailable.
        pass
    try:
        info = drive.files().get(
            fileId=folder, supportsAllDrives=True,
            fields="id,name,mimeType").execute()
    except Exception as exc:
        detail = drive_error_summary(exc)
        if oauth_email:
            detail += "; tài khoản OAuth đang dùng: " + oauth_email
        raise RuntimeError(detail) from exc
    if info.get("mimeType") != "application/vnd.google-apps.folder":
        raise RuntimeError("Mục đã cấu hình không phải thư mục Google Drive.")
    info["_oauthEmail"] = oauth_email
    st.session_state["_tdv_verified_photo_folder_id"] = folder
    return info


def configure_photo_folder():
    """ADMIN selects and verifies only the configured photo folder."""
    folder = str(st.secrets.get("photos_folder_id", "")).strip()
    oauth = st.secrets.get("google_drive_oauth", {})
    api_key = str(oauth.get("picker_api_key", "")).strip()
    if not folder:
        st.error("Thiếu photos_folder_id trong Streamlit Secrets.")
        return
    if not api_key:
        st.info("ADMIN cần thêm picker_api_key trong [google_drive_oauth] "
                "sau khi bật Google Picker API trên project OAuth.")
        return
    try:
        cred = drive_oauth_credentials()
    except Exception as exc:
        st.info("Chưa đủ OAuth Drive: " + str(exc))
        return
    st.caption("Google chỉ cấp quyền cho thư mục được chọn ở đây "
               "(phạm vi drive.file).")
    client_id = str(oauth.get("client_id", "")).strip()
    app_id = client_id.split("-", 1)[0] if "-" in client_id else ""
    if not app_id.isdigit():
        st.error("Client ID OAuth không chứa mã project hợp lệ để mở Picker.")
        return
    open_picker = st.button(
        "📂 MỞ CHỌN THƯ MỤC ẢNH",
        key="preview_open_photo_folder_picker",
        help="Mở cửa sổ Google Drive để chọn đúng thư mục ảnh DỰ ÁN."
    )
    picked = pick_folder(
        api_key, cred.token, app_id=app_id,
        key="preview_photo_folder_picker",
        open_picker=open_picker
    )
    if not picked:
        return
    if picked.get("folderId") != folder:
        st.error("Đã chọn nhầm thư mục. Hãy chọn đúng thư mục ảnh DỰ ÁN "
                 "đã cấu hình; chưa cấp quyền tải ảnh.")
        return
    try:
        # Verify with the exact OAuth credential/token Picker used.
        info = verify_configured_photo_folder(cred)
        st.success("Đã xác minh quyền OAuth trên đúng thư mục ảnh: " +
                   str(info.get("name") or folder))
    except Exception as exc:
        st.error("Google Picker đã chọn thư mục nhưng Drive chưa xác nhận quyền: " +
                 drive_error_summary(exc))


def drive_error_summary(error):
    """Return safe Google API diagnostics without exposing credentials."""
    response = getattr(error, "resp", None)
    status = getattr(response, "status", None)
    payload = {}
    try:
        raw = getattr(error, "content", b"")
        if isinstance(raw, bytes):
            raw = raw.decode("utf-8", errors="replace")
        payload = json.loads(raw) if raw else {}
    except (TypeError, ValueError):
        payload = {}
    detail = payload.get("error", {}) if isinstance(payload, dict) else {}
    reasons = []
    for item in detail.get("errors", []) if isinstance(detail, dict) else []:
        reason = item.get("reason")
        if reason and reason not in reasons:
            reasons.append(reason)
    message = detail.get("message", "") if isinstance(detail, dict) else ""
    if status is None:
        message = " ".join(str(error).split())
        return message[:240] if message else type(error).__name__
    label = "Google Drive HTTP " + str(status)
    if reasons:
        label += " (" + ", ".join(reasons[:3]) + ")"
    if message:
        label += ": " + " ".join(str(message).split())[:240]
    return label


@st.cache_data(ttl=45, show_spinner=False)
def read_sheet(name):
    sh, _ = client()
    try:
        values = _retry_google_read(lambda: sh.values_get("'" + name.replace("'", "''") + "'").get("values", []))
    except gspread.exceptions.APIError as error:
        raise PreviewSheetReadError(
            f"Không đọc được bảng {name}: Google Sheets HTTP {_google_api_status(error) or 'không rõ'}."
        ) from error
    if not values:
        return []
    headers = values[0]
    return [
        dict(zip(headers, row + [""] * (len(headers) - len(row))), _row=i)
        for i, row in enumerate(values[1:], 2) if any(str(x).strip() for x in row)
    ]

def sheet(name):
    return client()[0].worksheet(name)

def append(name, values):
    sheet(name).append_row([str(v) if v is not None else "" for v in values],
                           value_input_option="RAW")
    read_sheet.clear()

def audit(actor, action, target, target_id, result, detail=""):
    try:
        append("NHAT_KY_HE_THONG",
               [now(), actor, st.session_state.get("auth", {}).get("role", "VISITOR"),
                action, target, target_id, result, detail])
    except Exception:
        pass

def lookup(rows, key_field, key):
    return next((r for r in rows if r.get(key_field, "") == key), None)

def pbkdf(pin, salt):
    return hashlib.pbkdf2_hmac("sha256", pin.encode("utf-8"),
                              bytes.fromhex(salt), 390000).hex()

@st.cache_resource
def login_blocks():
    return {}


def effective_accounts():
    """Privileged credentials from Secrets only; team credentials from Sheet.

    Invalid/missing sheet fails CLOSED for sheet-backed team users,
    while existing ADMIN Secrets login remains available for recovery.
    """
    secrets_accounts = st.secrets.get("accounts", {})
    try:
        rows=read_sheet("QUAN_LY_DOI")
        accounts=team_accounts(secrets_accounts,
                               read_sheet("TAI_KHOAN_DOI"),rows)
        # Fail closed immediately for an externally deleted registration,
        # even before ADMIN reconciles site and account references.
        denied=revoked_team_ids(
            read_sheet("DANG_KY_THANH_VIEN"),
            read_sheet("LIEN_KET_DANG_KY_DOI"))
        return {name:info for name,info in accounts.items() if not (
            str(info.get("role","")).upper()=="DOI_TRUONG"
            and str(info.get("team_id","")) in denied)}
    except Exception:
        # Credential/source outage must never expose a deleted captain.
        # Keep privileged accounts in Secrets available for recovery.
        return {name:info for name,info in secrets_accounts.items()
                if str(info.get("role","")).upper() in ("ADMIN","LANH_DAO")}


def valid_login(username, pin):
    accounts = effective_accounts()
    info = accounts.get(username)
    if not info:
        return None
    locks = login_blocks()
    record = locks.get(username, {"count": 0, "until": 0})
    if record["until"] > time.time():
        return None
    try:
        valid = hmac.compare_digest(pbkdf(pin, info["salt"]), info["hash"])
    except (KeyError, ValueError):
        return None
    if valid:
        locks.pop(username, None)
        return {"username": username, "role": str(info["role"]).upper(),
                "team_id": str(info.get("team_id", "")),
                "_fp": account_fingerprint(info)}
    count = record["count"] + 1
    locks[username] = {"count": count, "until": time.time() + 900 if count >= 5 else 0}
    return None

def leadership_login():
    """Simple password-only entry for the read-only leadership dashboard."""
    st.subheader("📊 LÃNH ĐẠO – XEM BÁO CÁO")
    st.caption("ADMIN đặt một mật khẩu riêng trong Secrets. "
               "Nhập mật khẩu đó để xem số liệu và biểu đồ.")
    if st.button("← Trở về trang chính",key="leader_login_back"):
        st.session_state["leader_login_open"] = False
        st.rerun()

    configured_password = st.secrets.get("leader_view_password", "")
    if not isinstance(configured_password, str) or not configured_password:
        st.info('ADMIN thêm một dòng trong Secrets: '
                'leader_view_password = "mật khẩu tự đặt". '
                'Không dùng mật khẩu ADMIN.')
        return

    with st.form("leader_password_only",clear_on_submit=True):
        password = st.text_input("Mật khẩu Lãnh đạo",type="password")
        submitted = st.form_submit_button(
            "🔓 XEM BÁO CÁO",use_container_width=True)
    if not submitted:
        return
    if not check_leader_password(password, configured_password):
        st.error("Mật khẩu Lãnh đạo không đúng.")
        return

    # Keep only a password fingerprint in the session, never the password.
    st.session_state["auth"] = make_leader_session(configured_password)
    st.session_state["leader_login_open"] = False
    st.session_state["_remember_action"] = "remove"
    st.rerun()


# Stable PREVIEW URL for staff home-screen shortcuts (no production changes).
PREVIEW_HOME_URL = (
    "https://quan-ly-trien-khai-qfzumkrq7escxbxpv9v9tr.streamlit.app"
)


def home_screen_guide(expanded=False):
    """One-time device-owner shortcut setup; no silent install is possible."""
    with st.expander("📲 ĐƯA APP RA MÀN HÌNH ĐIỆN THOẠI",expanded=expanded):
        st.write("**Ghim logo TDV của QUẢN LÝ DỰ ÁN để mở app bằng một lần chạm.** "
                 "Nếu đã ghim icon cũ, hãy xóa icon đó rồi thêm lại từ app này.")
        st.caption("Mở app trong Safari hoặc Chrome. Nếu đang xem trong Zalo "
                   "hoặc Facebook, hãy chọn Mở bằng trình duyệt trước.")
        ios,android=st.tabs(["🍎 iPhone","🤖 Android"])
        with ios:
            st.markdown(
                "**①** Trong Safari: **Chia sẻ ⬆️ → Thêm vào Màn hình chính "
                "→ Thêm**. Bật *Mở dưới dạng ứng dụng web* nếu có."
            )
        with android:
            st.markdown(
                "**①** Trong Chrome: **⋮ → Cài đặt và tạo lối tắt** "
                "(hoặc **Thêm vào màn hình chính**) **→ Thêm**."
            )
        st.info("Biểu tượng app dùng logo TDV màu vàng cam. "
                "Mật khẩu tự đặt khi đăng ký chỉ dùng được **sau khi "
                "ADMIN duyệt**; có thể chọn ghi nhớ đăng nhập 30 ngày "
                "trên thiết bị cá nhân.")


def login(team_only=False):
    """Name-first team login or regular ADMIN login. PIN is always required."""
    if team_only:
        st.subheader("👷 ĐỘI TRƯỞNG – BÁO CÁO")
        st.caption("Nhập họ tên đã đăng ký và mật khẩu riêng. "
                   "Không cần biết mã đội hay mã tài khoản.")
        if st.button("← Trở về trang chính",key="team_login_back"):
            st.session_state["team_login_open"]=False
            st.rerun()
    login_teams=read_sheet("QUAN_LY_DOI")
    login_projects=sorted({str(t.get("Mã dự án","")).strip()
                           for t in login_teams
                           if t.get("Trạng thái")=="Hoạt động" and t.get("Mã dự án")})
    if team_only:
        # The project selection happens outside the form so changing project
        # immediately refreshes the name suggestions. Streamlit selectbox
        # supports built-in search while the member list is open.
        project_choice=st.selectbox(
            "Dự án của đội trưởng *",
            ["— Chọn dự án —"]+login_projects,
            key="team_login_project_first",
            help="Chọn dự án trước để chỉ tìm thành viên của dự án đó.")
        chosen_project=(project_choice
                        if project_choice!="— Chọn dự án —" else None)
        available_names=project_login_choices(
            login_teams,chosen_project,effective_accounts())
        typed=st.selectbox(
            "Họ và tên đội trưởng *",
            available_names,
            index=None,
            placeholder="Chạm và gõ vài chữ đầu của họ tên để tìm",
            disabled=not bool(chosen_project),
            key=f"team_login_member_{chosen_project or 'none'}",
            help="Chỉ hiện tên VC/KTV đang hoạt động và đã có tài khoản "
                 "thuộc dự án đã chọn.")
        if chosen_project and not available_names:
            st.info("Dự án này chưa có tài khoản VC/KTV hợp lệ để đăng nhập.")
    with st.form("team_login_form" if team_only else "login_form_preview"):
        if not team_only:
            # Keep the previous ADMIN/general login unchanged.
            typed=st.text_input("Họ tên đội trưởng hoặc mã tài khoản ADMIN")
            chosen_project=st.selectbox(
                "Dự án (chỉ chọn khi có đội trưởng trùng họ tên)",
                ["Tự xác định"]+login_projects,
                help="Nếu nhiều đội trưởng cùng tên, chọn đúng mã dự án. "
                     "Không cần chọn khi tên là duy nhất.")
        pin=st.text_input("Mật khẩu riêng",type="password")
        remember_key=remember_signing_key()
        remember=st.checkbox(
            "Ghi nhớ đăng nhập trên thiết bị này (30 ngày)",
            value=False,disabled=not bool(remember_key),
            help="Chỉ chọn trên điện thoại hoặc máy tính cá nhân; "
                 "không chọn trên thiết bị dùng chung.")
        if remember_key is None:
            st.caption("Ghi nhớ đăng nhập chỉ có khi ADMIN đã cấu hình khóa "
                       "ký phiên remember_login_secret trong Secrets.")
        submitted=st.form_submit_button(
            "✅ VÀO BÁO CÁO" if team_only else "ĐĂNG NHẬP",
            use_container_width=True)
    if not submitted:
        return
    if team_only and (not chosen_project or not typed):
        st.warning("Hãy chọn dự án và họ tên đội trưởng trước khi đăng nhập.")
        return
    accounts=effective_accounts()
    username,reason=resolve_login_username(
        typed,accounts,login_teams,
        project_id=None if chosen_project=="Tự xác định" else chosen_project)
    if reason=="ambiguous":
        st.error("Có nhiều hồ sơ đội trưởng trùng họ tên. "
                 "Vui lòng chọn đúng mã dự án ở phía trên; "
                 "nếu vẫn trùng tên, liên hệ ADMIN. Không tự chọn tài khoản khác.")
        return
    if not username:
        st.error("Họ tên hoặc tài khoản chưa được cấp quyền, hoặc thông tin "
                 "đăng nhập không hợp lệ. Vui lòng liên hệ ADMIN.")
        return
    if team_only and str(accounts[username].get("role","")).upper()!="DOI_TRUONG":
        st.error("Chỉ tài khoản đội trưởng được dùng nút này.")
        return
    user=valid_login(username,pin)
    if not user or (team_only and user["role"]!="DOI_TRUONG"):
        st.error("Mật khẩu không đúng hoặc tài khoản tạm khóa.")
        return
    st.session_state["auth"]=user
    st.session_state["team_login_open"]=False
    if remember and remember_key:
        account=accounts[user["username"]]
        st.session_state["_remember_token"]=issue_remember_token(
            user["username"],account,remember_key)
        st.session_state["_remember_action"]="set"
    else:
        st.session_state["_remember_action"]="remove"
    st.rerun()


def auth():
    a = st.session_state.get("auth")
    if not a:
        return None
    if a.get("username") == LEADER_SESSION_USERNAME:
        if verify_leader_session(
                a,st.secrets.get("leader_view_password","")):
            return a
        st.session_state.pop("auth",None)
        return None
    cur = effective_accounts().get(a.get("username"))
    if (not cur or str(cur.get("role","")).upper() != a.get("role")
            or str(cur.get("team_id","")) != a.get("team_id")
            or a.get("_fp") != account_fingerprint(cur)):
        st.session_state.pop("auth",None)
        return None
    if a["role"]=="DOI_TRUONG":
        match=[t for t in read_sheet("QUAN_LY_DOI")
               if t.get("Mã đội")==a.get("team_id") and
               t.get("Trạng thái")=="Hoạt động"]
        if len(match)!=1:
            st.session_state.pop("auth",None)
            return None
    return a

def require_role(*roles):
    a = auth()
    if not a or a["role"] not in roles:
        fail("Tài khoản không có quyền sử dụng chức năng này.")
    return a

def menu():
    sync_remembered_login()
    a = auth()
    if a and a.get("role")=="ADMIN":
        # ADMIN opens/refreshes the app: detect new registrations deleted
        # directly in Sheets, then revoke associated team and pending tasks.
        reconcile_deleted_captains()
    if not a:
        if st.session_state.get("leader_login_open"):
            return "__LEADERSHIP_LOGIN__"
        if st.session_state.get("team_login_open"):
            return "__TEAM_LOGIN__"
        if st.button("📊 LÃNH ĐẠO – XEM BÁO CÁO",
                     type="primary",use_container_width=True,
                     key="leader_entry_button"):
            st.session_state["leader_login_open"]=True
            st.rerun()
        if st.button("👷 ĐỘI TRƯỞNG – BÁO CÁO",
                     use_container_width=True,key="team_entry_button"):
            st.session_state["team_login_open"]=True
            st.rerun()
        if not st.session_state.get("registered_notice"):
            home_screen_guide(expanded=False)
        opts = ["🔐 Đăng nhập ADMIN / tài khoản khác",
                "📝 Đăng ký tham gia"]
    elif a["role"] == "ADMIN":
        opts = ["📊 Báo cáo tiến độ", "📁 Dự án/Điểm/Nhập kho",
                "🧹 RESET / XÓA DỰ ÁN", "📦 Phân bổ",
                "👥 Đăng ký & quản lý đội", "📦 Bàn giao tập kết",
                "📝 Đăng ký hộ đội trưởng", "📋 Báo cáo nhiệm vụ",
                "🗺️ Tìm đường"]
    elif a["role"] == "LANH_DAO":
        # Executive landing is always the read-only progress view.
        st.caption("📊 LÃNH ĐẠO • Chế độ chỉ xem")
        if st.button("🚪 ĐĂNG XUẤT",use_container_width=True,
                     key="leader_logout"):
            st.session_state.pop("auth",None)
            st.session_state["leader_login_open"]=False
            st.session_state["_remember_action"]="remove"
            st.rerun()
        return "📊 Báo cáo tiến độ"
    elif a["role"] == "DOI_TRUONG":
        opts = ["📋 Báo cáo nhiệm vụ", "🗺️ Tìm đường"]
    else:
        opts = ["🗺️ Tìm đường"]
    if a:
        name=friendly_account_name(a,read_sheet("QUAN_LY_DOI")) if a["role"]=="DOI_TRUONG" else a["username"]
        st.caption(f"Đã đăng nhập: {name} • {a['role']}")
        if st.button("Đăng xuất"):
            st.session_state.pop("auth", None)
            st.session_state["_remember_action"] = "remove"
            st.rerun()
    if a and a["role"] == "ADMIN":
        pending_count=sum(x.get("Trạng thái duyệt")=="Chờ duyệt"
                          for x in read_sheet("DANG_KY_THANH_VIEN"))
        if pending_count:
            st.info(f"📩 Có {pending_count} hồ sơ đội trưởng CHỜ DUYỆT.")
            if st.button("📩 MỞ DANH SÁCH CHỜ DUYỆT", use_container_width=True):
                st.session_state["preview_menu"] = "👥 Đăng ký & quản lý đội"
                st.rerun()
    # Navigation requests are applied before widget construction on the
    # next run; writing to a rendered selectbox's key is not permitted.
    destination=st.session_state.pop("_next_menu",None)
    if destination in opts:
        st.session_state["preview_menu"]=destination
    project_to_assign=st.session_state.pop("_next_assignment_project",None)
    if project_to_assign and destination=="👥 Đăng ký & quản lý đội":
        st.session_state["site_assign_project"]=project_to_assign
    return st.selectbox("CHỌN CHỨC NĂNG", opts, key="preview_menu")

@st.cache_resource
def registration_submit_lock():
    """One lock per app process for signup checks and writes."""
    return threading.Lock()


def register(on_behalf=False):
    """Public registration and Admin-assisted registration use the same pending queue."""
    acting_admin = require_role("ADMIN") if on_behalf else None
    st.subheader("📝 " + ("ADMIN đăng ký hộ đội trưởng" if on_behalf else "Đội trưởng đăng ký tham gia dự án"))
    st.caption("Khai khu vực đã hợp đồng với công ty. Đây là thông tin đăng ký, "
               "không tự cấp quyền báo cáo; ADMIN sẽ xét duyệt và giao nhiệm vụ.")
    saved_notice = st.session_state.pop("registered_notice", None)
    if saved_notice:
        st.success(saved_notice)
        if not on_behalf:
            st.link_button("📖 XEM HƯỚNG DẪN & THỰC HÀNH THỬ",
                           GUIDE_URL,use_container_width=True)
            home_screen_guide(expanded=True)
    projects = [p for p in read_sheet("DANH_SACH_DU_AN")
                if p.get("Mã dự án") and p.get("Tên dự án")]
    if not projects:
        st.info("Chưa có dự án mở đăng ký. Vui lòng liên hệ ADMIN.")
        return
    options = [p["Mã dự án"] for p in projects]
    project_names = {p["Mã dự án"]: p["Tên dự án"] for p in projects}
    points = read_sheet("DANH_SACH_DIEM")
    revision = st.session_state.get("registration_form_revision", 0)
    # Project selector stays OUTSIDE the form: changing project must immediately
    # refresh the valid list of locations, never silently retain another project.
    project = st.selectbox("Dự án đăng ký *", options,
        format_func=lambda key: f"{key} — {project_names[key]}",
        key=f"register_project_{'admin' if on_behalf else 'public'}")
    with st.form(f"registration_{'admin' if on_behalf else 'public'}_{revision}"):
        full_name = st.text_input("Họ và tên *", max_chars=120)
        phone = st.text_input("Điện thoại liên hệ (Zalo) *", max_chars=20)
        area_options = sorted({
            p["Tên xã/phường/đơn vị"] for p in points
            if p.get("Mã dự án") == project and p.get("Tên xã/phường/đơn vị")
        })
        area = st.multiselect("Khu vực đã hợp đồng (chọn nhiều nếu cần)",
                              area_options, key=f"register_areas_{project}_{revision}")
        area_other = st.text_area("Khu vực theo hợp đồng cần ghi bổ sung",
            placeholder="Gõ địa bàn, xã/phường hoặc đơn vị chưa có trong danh sách...",
            max_chars=1200, height=84)
        specialty = st.selectbox("Chuyên môn / Nhiệm vụ *",
            ["Vận chuyển / Giao nhận", "KTV lắp đặt",
             "Vận chuyển & Lắp đặt", "Điều phối / Giám sát", "Kho vận / Hậu cứ"])
        vehicle = st.selectbox("Phương tiện", ["Xe máy", "Xe tải", "Xe bán tải", "Khác"])
        if not on_behalf:
            st.markdown("#### 🔐 Tự đặt mật khẩu đăng nhập")
            st.caption("V tự chọn mật khẩu dễ nhớ, **từ 6 ký tự có chữ và số**, "
                       "ví dụ một cụm từ riêng có số (không dùng ví dụ chung). "
                       "Không dùng họ tên, số điện thoại, ngày sinh hay "
                       "mật khẩu Zalo. Chỉ lưu mã băm; "
                       "ADMIN duyệt xong mới đăng nhập được.")
            desired_password=st.text_input(
                "Mật khẩu tự chọn *",type="password",max_chars=128)
            confirm_password=st.text_input(
                "Nhập lại mật khẩu *",type="password",max_chars=128)
        st.caption("Khu vực tự khai chỉ để ADMIN tham khảo khi phân công. "
                   "Chỉ điểm được giao chính thức mới có thể báo cáo.")
        if st.form_submit_button("GỬI ĐĂNG KÝ" if not on_behalf else
                                 "GỬI HỒ SƠ ĐĂNG KÝ HỘ", use_container_width=True):
            normalized_phone = "".join(c for c in phone if c.isdigit())
            described_areas = ", ".join(area + ([area_other.strip()] if area_other.strip() else []))
            if len(full_name.strip()) < 3 or len(normalized_phone) not in (9,10,11):
                st.error("Vui lòng kiểm tra họ tên và số điện thoại (9–11 chữ số).")
            elif not on_behalf and desired_password!=confirm_password:
                st.error("Hai lần nhập mật khẩu chưa giống nhau.")
            elif not on_behalf and validate_custom_password(
                    desired_password,full_name,normalized_phone):
                st.error(validate_custom_password(
                    desired_password,full_name,normalized_phone))
            elif not described_areas:
                st.error("Hãy chọn hoặc ghi ít nhất một khu vực đã hợp đồng.")
            elif project not in [p["Mã dự án"] for p in read_sheet("DANH_SACH_DU_AN")]:
                st.error("Dự án không còn tồn tại, vui lòng kiểm tra lại.")
            else:
                with registration_submit_lock():
                    # Prevent concurrent duplicate signups; bypass cached rows.
                    read_sheet.clear()
                    existing = read_sheet("DANG_KY_THANH_VIEN")
                    same = [row for row in existing
                        if row.get("Mã dự án đăng ký") == project
                        and "".join(c for c in row.get("Số điện thoại","") if c.isdigit()) == normalized_phone
                        and " ".join(row.get("Họ tên","").casefold().split()) ==
                            " ".join(full_name.casefold().split())
                        and row.get("Trạng thái duyệt") in ("Chờ duyệt", "Đã duyệt")]
                    if same:
                        st.warning("Hồ sơ này đã đăng ký trước đó. "
                                   "Không được đăng ký lại để đổi mật khẩu. "
                                   "Nếu muốn đổi mật khẩu tự chọn hoặc hồ sơ cũ "
                                   "chưa có mật khẩu, hãy nhờ ADMIN vào mục "
                                   "CẤP / ĐẶT LẠI MẬT KHẨU ĐỘI TRƯỞNG "
                                   "và nhập mật khẩu theo đề nghị.")
                    else:
                        rid = "DK-" + uuid.uuid4().hex[:12].upper()
                        submitter = ("ADMIN:" + acting_admin["username"]) if acting_admin else "Tự đăng ký"
                        try:
                            date_serial = sheet_datetime_serial()
                            registration=[rid, date_serial, full_name.strip(),
                                          normalized_phone,described_areas,
                                          specialty,vehicle,"Chờ duyệt",
                                          "","","",project,submitter]
                            if on_behalf:
                                sheet("DANG_KY_THANH_VIEN").append_row(
                                    registration,value_input_option="RAW")
                                notice="ADMIN cần cấp mật khẩu riêng sau khi duyệt."
                            else:
                                pending=pending_row(rid,project,desired_password,now())
                                # One atomic Sheets request: no registration without
                                # its corresponding salted password hash.
                                batch=registration_requests(
                                    sheet("DANG_KY_THANH_VIEN").id,
                                    sheet("MAT_KHAU_DANG_KY").id,
                                    registration,pending)
                                client()[0].batch_update({"requests":batch})
                                notice=("Mật khẩu đã được lưu an toàn dưới dạng mã "
                                        "băm. Chỉ đăng nhập được sau khi ADMIN duyệt.")
                            read_sheet.clear()
                            st.session_state["registered_notice"] = (
                                f"Đã nhận hồ sơ {rid} — đang chờ ADMIN duyệt. "+notice)
                            st.session_state["registration_form_revision"] = revision + 1
                            st.rerun()
                        except Exception as e:
                            st.error("Chưa xác nhận gửi hồ sơ thành công (" +
                                     type(e).__name__ +
                                     "). Hãy nhờ ADMIN kiểm tra danh sách trước khi gửi lại.")

def directions():
    st.subheader("🗺️ Tìm địa điểm – Dẫn đường")
    points = [r for r in read_sheet("DANH_SACH_DIEM") if r.get("Mã điểm") and r.get("Tên xã/phường/đơn vị")]
    if not points:
        st.info("Admin chưa tạo danh sách địa điểm.")
        return
    options = {f"{p['Tên xã/phường/đơn vị']} • {p['Mã điểm']}": p for p in points}
    selected = st.selectbox("Gõ tên xã, phường hoặc đơn vị cần đến", list(options))
    p = options[selected]
    st.caption(p.get("Địa chỉ chi tiết") or p.get("Tỉnh/Thành") or "")
    st.info("Có thể cho phép GPS để dùng vị trí đang đứng, hoặc để Google Maps tự xác định điểm xuất phát.")
    loc = streamlit_geolocation()
    lat = loc.get("latitude") if isinstance(loc, dict) else None
    lon = loc.get("longitude") if isinstance(loc, dict) else None
    destination = ((p.get("Vĩ độ") + "," + p.get("Kinh độ"))
                   if p.get("Vĩ độ") and p.get("Kinh độ")
                   else ", ".join(filter(None, [p["Tên xã/phường/đơn vị"],
                                                 p.get("Địa chỉ chi tiết"), p.get("Tỉnh/Thành")])))
    url = "https://www.google.com/maps/dir/?api=1&destination=" + quote(destination)
    if lat is not None and lon is not None:
        url += "&origin=" + quote(f"{lat},{lon}")
    if not (p.get("Vĩ độ") and p.get("Kinh độ")):
        st.warning("Điểm này chưa có GPS đã xác minh. Google Maps sẽ tìm "
                   "theo tên/địa chỉ; hãy chọn đúng vị trí đích trên bản đồ "
                   "trước khi bắt đầu chỉ đường.")
    st.link_button("🧭 MỞ GOOGLE MAPS CHỈ ĐƯỜNG", url, use_container_width=True)
    st.caption("Tìm đường không cấp quyền báo cáo nhiệm vụ.")


def bulk_import_locations():
    """Admin-only, review-before-write import of a project location list."""
    require_role("ADMIN")
    st.markdown("#### 📥 Nạp danh sách xã, đơn vị, địa điểm")
    st.caption("Dữ liệu chỉ nhập vào dự án đã chọn. Không xóa hoặc ghi đè các điểm đã có.")
    projects = [r["Mã dự án"] for r in read_sheet("DANH_SACH_DU_AN") if r.get("Mã dự án")]
    if not projects:
        st.info("Admin cần tạo dự án trước khi nhập danh sách.")
        return
    project = st.selectbox("Dự án cần nhập danh sách *", projects, key="bulk_location_project")
    st.download_button("⬇️ Tải mẫu Excel (1 cột tên địa điểm)",
        data=blank_excel_template(),
        file_name="MAU_DANH_SACH_DIEM.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="download_location_template")
    upload = st.file_uploader("Tải Word (.docx) hoặc Excel (.xlsx), tối đa 5 MB",
        type=["docx", "xlsx"], key="upload_bulk_locations")
    if upload is None:
        return
    try:
        sources = read_location_file(upload.name, upload.getvalue())
    except Exception as exc:
        st.error("Không đọc được danh sách: " + str(exc))
        return
    if not sources:
        st.warning("Tệp không có dòng dữ liệu hợp lệ.")
        return
    chosen_sheet = st.selectbox("Trang chứa danh sách", list(sources),
                                key="bulk_location_source_sheet")
    rows = sources[chosen_sheet]
    number_of_cols = min(24, max((len(row) for row in rows), default=1))
    preferred_col = min(column_choice(rows), number_of_cols - 1)
    cols = list(range(number_of_cols))
    selected_col = st.selectbox("Cột chứa tên xã/phường/đơn vị", cols,
        index=preferred_col, key="bulk_location_column",
        format_func=lambda idx: f"Cột {idx+1} • " +
        ", ".join(str(r[idx])[:48] for r in rows[:4] if len(r)>idx and r[idx])[:125])
    first_is_header = bool(rows and len(rows[0]) > selected_col and is_header(rows[0][selected_col]))
    skip_header = st.checkbox("Bỏ qua dòng đầu (dòng tiêu đề)",
                              value=first_is_header, key="bulk_location_skip_header")
    names = candidate_names(rows, selected_col, skip_header)
    seen, unique = set(), []
    duplicate_inside = 0
    for name in names:
        norm = normalize_name(name)
        if norm in seen:
            duplicate_inside += 1
        else:
            seen.add(norm)
            unique.append(name)
    prior = [r for r in read_sheet("DANH_SACH_DIEM") if r.get("Mã dự án") == project]
    oldnames = {normalize_name(r.get("Tên xã/phường/đơn vị")) for r in prior}
    to_add = [name for name in unique if normalize_name(name) not in oldnames]
    st.write(f"**Đọc được:** {len(names)} dòng  •  **Trùng trong tệp:** {duplicate_inside}"
             f"  •  **Đã có trong dự án:** {len(unique)-len(to_add)}"
             f"  •  **Dự kiến thêm:** {len(to_add)}")
    if to_add:
        display_centered_table([{"STT":i, "Tên xã/phường/đơn vị":name}
                      for i,name in enumerate(to_add[:100],1)],
                     hide_index=True, use_container_width=True)
        if len(to_add)>100:
            st.caption("Đang hiển thị 100 tên đầu tiên; khi xác nhận sẽ nhập cả danh sách.")
    else:
        st.info("Không có địa điểm mới để thêm.")
        return
    st.warning("Vui lòng kiểm tra cột, tên địa điểm và dự án trước khi nạp.")
    if st.button(f"✅ XÁC NHẬN NẠP {len(to_add)} ĐỊA ĐIỂM",
                 type="primary", use_container_width=True,
                 key="confirm_bulk_locations", disabled=len(to_add)>1000):
        # Recheck project and duplicates immediately before committing.
        if project not in {r["Mã dự án"] for r in read_sheet("DANH_SACH_DU_AN")}:
            st.error("Dự án không còn tồn tại; không ghi dữ liệu.")
            return
        latest = [r for r in read_sheet("DANH_SACH_DIEM") if r.get("Mã dự án") == project]
        already = {normalize_name(r.get("Tên xã/phường/đơn vị")) for r in latest}
        final = [name for name in to_add if normalize_name(name) not in already]
        if not final:
            st.info("Danh sách này đã được nạp trước đó, không tạo bản ghi trùng.")
            return
        payload = []
        for name in final:
            point_id = "DD-" + uuid.uuid4().hex[:12].upper()
            # Coordinates must be verified by Admin before relying on navigation.
            payload.append([point_id, project, name, "", "", "", "", "", "",
                            "Hoạt động", "Nhập từ " + upload.name[:80]])
        try:
            sheet("DANH_SACH_DIEM").append_rows(payload, value_input_option="RAW")
            read_sheet.clear()
            audit(auth()["username"], "NHẬP ĐỊA ĐIỂM", "DANH_SACH_DIEM", project,
                  "OK", f"{len(final)} điểm từ tệp {upload.name[:80]}")
            st.success(f"Đã thêm {len(final)} địa điểm vào dự án {project}. "
                       "Điểm trùng được giữ nguyên, không ghi đè.")
            st.info("✅ Danh sách đã nạp thành công. Chưa có tọa độ GPS "
                    "không ảnh hưởng đến nhập kho, PHÂN BỔ hay GÁN ĐỘI. "
                    "Khi cần dẫn đường, Google Maps có thể tìm theo tên điểm "
                    "nhưng ADMIN/đội trưởng phải kiểm tra đúng vị trí trước "
                    "khi đi; không tự coi kết quả tìm tên là tọa độ xác nhận.")
        except Exception as exc:
            read_sheet.clear()
            st.error("Chưa thể xác nhận nhập danh sách thành công (" +
                     type(exc).__name__ + "). Hãy kiểm tra Google Sheets trước khi thử lại.")

def projects_and_inventory():
    require_role("ADMIN")
    st.subheader("📁 Danh mục do Admin quản lý")
    choice = st.radio("Thao tác quản lý", ["Dự án", "Địa điểm", "Hàng nhập kho", "Nạp danh sách địa điểm"], horizontal=False)
    if choice == "Nạp danh sách địa điểm":
        bulk_import_locations()
        return
    projects = read_sheet("DANH_SACH_DU_AN")
    ids = [x.get("Mã dự án") for x in projects if x.get("Mã dự án")]
    # This selector MUST remain OUTSIDE clear_on_submit form.
    # Otherwise after every receipt it silently reverts to first project.
    project = None
    if choice in ("Địa điểm", "Hàng nhập kho"):
        if not ids:
            st.warning("Chưa có dự án. Hãy tạo dự án trước.")
            return
        project = st.selectbox(
            "📁 Dự án đang nhập (giữ nguyên sau mỗi lần lưu)",
            ids, key="master_sticky_project")
        st.caption(f"Đang ghi dữ liệu vào dự án: {project}. "
                   "Khi nhập phiếu mới, kiểm tra mã dự án trước khi bấm LƯU.")
    with st.form("new_master", clear_on_submit=True):
        if choice == "Dự án":
            code = st.text_input("Mã dự án")
            name = st.text_input("Tên dự án")
            deadline = st.date_input("Hạn hoàn thành")
        elif choice == "Địa điểm":
            code = st.text_input("Mã điểm duy nhất")
            name = st.text_input("Tên xã/phường/đơn vị")
            address = st.text_input("Địa chỉ chi tiết")
            province = st.text_input("Tỉnh/Thành")
            route = st.text_input("Tuyến/Khu vực")
            latitude = st.text_input("Vĩ độ (nếu có)")
            longitude = st.text_input("Kinh độ (nếu có)")
        else:
            sku = st.text_input("SKU")
            name = st.text_input("Tên hàng hóa")
            unit = st.text_input("ĐVT")
            quantity = st.number_input("Số lượng THỰC NHẬN", min_value=0.0, step=1.0)
        if st.form_submit_button("LƯU", use_container_width=True):
            try:
                if choice == "Dự án":
                    if not code or not name or code in ids:
                        fail("Trùng hoặc thiếu mã dự án/tên dự án.")
                    append("DANH_SACH_DU_AN",[code,name,now()[:10],deadline.isoformat(),"Đang triển khai",auth()["username"],""])
                elif choice == "Địa điểm":
                    if not code or not name or lookup(read_sheet("DANH_SACH_DIEM"),"Mã điểm",code):
                        fail("Trùng hoặc thiếu mã điểm/tên điểm.")
                    url="https://www.google.com/maps/search/?api=1&query="+quote(", ".join(filter(None,[name,address,province])))
                    append("DANH_SACH_DIEM",[code,project,name,address,province,route,latitude,longitude,url,"Hoạt động",""])
                else:
                    if not sku or not name or not unit or quantity <= 0:
                        fail("Thiếu thông tin hoặc SLHH không hợp lệ.")
                    append("NHAP_KHO",["NK-"+uuid.uuid4().hex[:12].upper(),project,now(),sku,name,unit,quantity,"",auth()["username"],""])
                audit(auth()["username"],"THÊM DANH MỤC",choice,code if choice != "Hàng nhập kho" else sku,"OK")
                read_sheet.clear()
                st.success("Đã lưu vào Google Sheets DỰ ÁN.")
            except Exception as e:
                st.error(f"Lưu thất bại: {type(e).__name__}")
    if choice == "Hàng nhập kho":
        st.markdown("#### 📦 Hàng còn chờ phân bổ theo SKU")
        st.caption("Tổng thực nhập qua nhiều đợt trừ tổng đã phân bổ cho "
                   "địa điểm. Đây KHÔNG phải tồn vật lý sau xuất kho.")
        records,problems = allocation_waiting_summary(
            project,read_sheet("NHAP_KHO"),read_sheet("KHO_PHAN_BO"))
        if problems:
            for problem in problems[:8]:
                st.error(problem)
            st.warning("Dữ liệu kho có lỗi; cần đối soát trước khi phân bổ.")
        elif records:
            display_centered_table(records,hide_index=True,use_container_width=True)
            if all(row["Hàng còn chờ phân bổ"] == 0 for row in records):
                st.success("Toàn bộ hàng đã nhập của dự án đã được phân bổ hết "
                           "(còn chờ phân bổ = 0 cho mọi SKU).")
        else:
            st.info("Dự án chưa nhập hàng. Khi có phiếu nhập, "
                    "tổng hàng và số còn chờ phân bổ sẽ hiển thị tại đây.")

def read_sheet_live(name):
    """Bypass Streamlit's short cache for the final stock/duplicate check."""
    values = _retry_google_read(lambda: client()[0].values_get("'" + name.replace("'", "''") + "'").get("values", []))
    if not values:
        return []
    headers = values[0]
    return [dict(zip(headers, row + [""] * (len(headers)-len(row))), _row=i)
            for i, row in enumerate(values[1:], 2) if any(str(x).strip() for x in row)]


def reconcile_deleted_captains():
    """ADMIN-triggered automatic cleanup of externally deleted profiles.

    No background Google Sheets trigger exists: runs when logged-in ADMIN
    opens or refreshes the PREVIEW app, without a separate deletion button.
    Legacy/manual teams not linked in LIEN_KET_DANG_KY_DOI are untouched.
    """
    who=st.session_state.get("auth",{})
    if who.get("role")!="ADMIN":
        return
    try:
        # No external reads/writes until at least one new tracked link exists.
        links=read_sheet("LIEN_KET_DANG_KY_DOI")
        if not any(r.get("Trạng thái liên kết")==LINK_ACTIVE for r in links):
            return
        fresh_links=read_sheet_live("LIEN_KET_DANG_KY_DOI")
        fresh_regs=read_sheet_live("DANG_KY_THANH_VIEN")
        orphaned=lost_registrations(fresh_regs,fresh_links)
        if not orphaned:
            return
        total_tasks=0
        reviewed=0
        processed=0
        for link in orphaned:
            # Each linked team is reconciled against fresh data before
            # applying a single atomic batch across the related worksheets.
            teams=read_sheet_live("QUAN_LY_DOI")
            creds=read_sheet_live("TAI_KHOAN_DOI")
            tasks_vc=read_sheet_live("VAN_CHUYEN")
            tasks_ld=read_sheet_live("LAP_DAT")
            reports=read_sheet_live("BAO_CAO_TRIEN_KHAI")
            hub=read_sheet_live("BAN_GIAO_TAP_KET")
            plan=plan_detach(link,teams,creds,tasks_vc,tasks_ld,reports,hub)
            keys=("QUAN_LY_DOI","TAI_KHOAN_DOI","VAN_CHUYEN",
                  "LAP_DAT","LIEN_KET_DANG_KY_DOI")
            ids={n:sheet(n).id for n in keys}
            changes=detach_requests(plan,ids,now())
            client()[0].batch_update({"requests":changes})
            read_sheet.clear()
            processed+=1
            total_tasks+=len(plan["pending_tasks"])
            reviewed+=int(plan["review"])
            audit(who["username"],"TỰ GỠ ĐỘI SAU XÓA ĐĂNG KÝ",
                  "LIEN_KET_DANG_KY_DOI",
                  str(link.get("Mã đăng ký","")),"OK",
                  f"Đội: {plan['team_id']} • Nhiệm vụ trả về: "
                  f"{len(plan['pending_tasks'])} • Cần đối soát: {plan['review']}")
        if processed:
            st.success(f"Đã đối chiếu hồ sơ bị xóa: gỡ {processed} đội "
                       f"và trả {total_tasks} nhiệm vụ/SKU chưa thực hiện "
                       "về chờ phân công lại.")
            if reviewed:
                st.warning(f"{reviewed} đội có báo cáo, chứng từ tập kết "
                           "hoặc nhiệm vụ đã thực hiện: đã khóa quyền, "
                           "giữ chứng từ cũ để ADMIN đối soát.")
    except Exception as exc:
        read_sheet.clear()
        st.warning("Không hoàn tất được đồng bộ xóa đăng ký ("+
                   type(exc).__name__+"). Không xóa thủ công nhiệm vụ, "
                   "hãy vào Quản lý đội để đối soát. "
                   "Các tài khoản thuộc hồ sơ đã mất vẫn bị khóa "
                   "khi xác thực lại.")


def bulk_matrix_xlsx(plan, sku_list):
    wb = ExcelWorkbook()
    ws = wb.active
    ws.title = "PHAN_BO_XEM_TRUOC"
    ws.append(["Mã điểm", "Tên địa điểm"] + sku_list)
    for r in plan:
        ws.append([r["Mã điểm"], r["Tên địa điểm"]] +
                  [float(r["quantities"][sku]) for sku in sku_list])
    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 44
    for n in range(len(sku_list)):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(n+3)].width = 16
    for cell in ws[1]:
        from openpyxl.styles import Font, PatternFill
        cell.fill = PatternFill("solid", fgColor="183A72")
        cell.font = Font(color="FFFFFF", bold=True)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def allocate_bulk():
    require_role("ADMIN")
    st.subheader("📦 Phân bổ theo nhóm địa điểm – bộ định mức hồ sơ thầu")
    st.caption("Khai một bộ SKU và số lượng cấp cho MỖI ĐIỂM, sau đó chọn "
               "nhiều xã/đơn vị cùng nhận bộ hàng đó. Nhóm có định mức khác "
               "thực hiện ở đợt riêng hoặc chỉnh ngoại lệ. "
               "PHÂN BỔ không phải XUẤT KHO.")
    prows=read_sheet("DANH_SACH_DU_AN")
    projects=[r["Mã dự án"] for r in prows if r.get("Mã dự án")]
    if not projects:
        st.warning("Cần tạo dự án trước.")
        return
    project=st.selectbox("Dự án",projects,key="bulk_project")
    points=[r for r in read_sheet("DANH_SACH_DIEM")
            if r.get("Mã dự án")==project and r.get("Mã điểm")
            and r.get("Trạng thái")=="Hoạt động"]
    if len(points)!=len({r["Mã điểm"] for r in points}):
        st.error("Danh sách địa điểm có mã trùng. Cần ADMIN đối soát trước.")
        return
    imports=read_sheet("NHAP_KHO")
    allocations=read_sheet("KHO_PHAN_BO")
    inventory,inventory_issues=inventory_by_sku(project,imports,allocations)
    if inventory_issues:
        for issue in inventory_issues:
            st.error(issue)
        st.warning("Phải đối soát kho trước khi phân bổ.")
        return
    if not points or not inventory:
        st.warning("Dự án cần có danh sách địa điểm và ít nhất một SKU trong NHAP_KHO.")
        return
    point_by_id={r["Mã điểm"]:r for r in points}
    point_ids=list(point_by_id)
    names={pid:point_by_id[pid]["Tên xã/phường/đơn vị"] for pid in point_ids}
    skus=list(inventory)
    st.write(f"**{len(point_ids)} địa điểm** · **{len(skus)} SKU nhập kho** "
             "· Có thể chia thành nhiều nhóm định mức theo hồ sơ thầu.")

    st.markdown("#### 1. Bộ mặt hàng và định mức cho mỗi điểm")
    st.caption("Nhập số lượng **mỗi địa điểm được cấp**, KHÔNG phải tổng "
               "của cả nhóm. Ví dụ 12 bộ/điểm × 20 điểm = 240 bộ. "
               "SKU nào không cấp trong nhóm này thì để 0.")
    defaults={}
    for sku in skus:
        item=inventory[sku]
        st.caption(f"**{sku}** — {item['name']} ({item['unit']}) · "
                   f"Kho còn khả dụng để phân bổ: {item['remaining']:g}")
        defaults[sku]=st.number_input(
            f"Số lượng {sku} / mỗi điểm ({item['unit']})",
            min_value=0.0,value=0.0,step=1.0,format="%.3f",
            key=f"bulk_default_{project}_{sku}")

    st.markdown("#### 2. Chọn nhóm xã/đơn vị có cùng định mức")
    method=st.radio("Cách chọn danh sách điểm theo hồ sơ thầu",
        ["🔎 Gõ tìm tên / mã điểm", "📋 Dán danh sách từ hợp đồng",
         "📄 Nạp danh sách Excel", "☑️ Chọn tất cả điểm chưa phân bổ"],
        key=f"bulk_method_{project}")
    issues=[]
    if method.startswith("🔎"):
        chosen_ids=st.multiselect(
            "Địa điểm nhận bộ hàng hóa trên",
            point_ids, default=[],
            format_func=lambda pid:f"{names[pid]} • {pid}",
            placeholder="Gõ tên xã, đơn vị hoặc mã điểm để tìm...",
            key=f"bulk_points_{project}")
        st.caption("Chọn liên tiếp các điểm trong cùng một nhóm. "
                   "Không phải kéo qua cả 136 hay 500 điểm.")
    elif method.startswith("📋"):
        text_list=st.text_area(
            "Dán danh sách địa điểm đã ký hợp đồng",
            placeholder="Mỗi dòng một tên xã/đơn vị hoặc một Mã điểm. "
                        "Ví dụ:\nPhường Minh Xuân\nDD-TEST-127",
            key=f"bulk_contract_text_{project}",height=170)
        chosen_ids,issues=match_contract_locations(
            project,points,text_list.splitlines())
    elif method.startswith("📄"):
        upload=st.file_uploader(
            "Nạp Excel danh sách điểm cùng định mức "
            "(cột Mã điểm hoặc Tên xã/phường/đơn vị)",
            type=["xlsx"],key=f"bulk_group_excel_{project}")
        chosen_ids=[]
        if upload:
            try:
                raw_names=read_contract_locations_excel(
                    upload.name,upload.getvalue())
                chosen_ids,issues=match_contract_locations(
                    project,points,raw_names)
            except ValueError as exc:
                issues=[str(exc)]
    else:
        has_any={r.get("Mã điểm") for r in allocations
                 if r.get("Mã dự án")==project and r.get("SKU")}
        chosen_ids=[pid for pid in point_ids if pid not in has_any]
        st.caption("Chỉ chọn điểm chưa có SKU nào được phân bổ. "
                   "Điểm đã phân bổ một phần phải kiểm tra riêng.")

    for issue in issues[:10]:
        st.error(issue)
    if len(issues)>10:
        st.error(f"Còn {len(issues)-10} địa điểm không khớp. "
                 "Hãy sửa danh sách, không thể ghi.")
    st.write(f"**Nhóm hiện chọn: {len(chosen_ids)} địa điểm** "
             f"trong tổng số {len(point_ids)} điểm của dự án.")
    if chosen_ids:
        with st.expander("Xem tên và mã điểm đã chọn"):
            display_centered_table([{"Mã điểm":pid,"Xã/phường/đơn vị":names[pid]}
                          for pid in chosen_ids],
                         hide_index=True,use_container_width=True)
    if issues:
        return
    if not chosen_ids:
        st.info("Vui lòng chọn ít nhất một địa điểm để lập phương án phân bổ.")
        return

    # Basket totals here are a preview, not a record of a warehouse issue.
    basket_totals=preview_basket(skus,defaults,len(chosen_ids))
    st.markdown("#### Đối chiếu nhanh bộ định mức × số điểm")
    display_centered_table([{
        "SKU":sku,"Hàng hóa":inventory[sku]["name"],
        "ĐVT":inventory[sku]["unit"],
        "SL/điểm":float(defaults[sku]),
        "Số điểm":len(chosen_ids),
        "Tổng dự kiến":float(basket_totals[sku]),
        "Tồn có thể phân bổ":float(inventory[sku]["remaining"])
    } for sku in skus if basket_totals[sku]>0],
        hide_index=True,use_container_width=True)
    st.caption("Đây là bảng rà soát theo định mức ADMIN vừa nhập; "
               "chưa có dữ liệu định mức thầu độc lập để tự so sánh với "
               "hồ sơ gốc. ADMIN cần kiểm tra định mức trước khi xác nhận.")

    st.markdown("#### 3. Điều chỉnh những đơn vị có định mức khác")
    uploaded = st.file_uploader(
        "Tùy chọn: nạp Excel các điểm ngoại lệ (cột Mã điểm và các cột SKU)",
        type=["xlsx"], key=f"bulk_exc_upload_{project}")
    imported, imported_points = {}, []
    upload_fingerprint = "noexcel"
    if uploaded is not None:
        try:
            import_data = uploaded.getvalue()
            imported, imported_points = parse_exception_excel(
                uploaded.name, import_data, chosen_ids, skus)
            upload_fingerprint = hashlib.sha256(import_data).hexdigest()[:10]
            st.success(f"Đã đọc {len(imported_points)} điểm ngoại lệ trong Excel. "
                       "Chưa ghi vào Google Sheets.")
        except Exception as exc:
            st.error("Excel ngoại lệ không hợp lệ: " + str(exc))
            return
    exceptions = st.multiselect("Chọn các điểm ngoại lệ (có thể nhiều hoặc ít hơn 15)",
        chosen_ids,default=imported_points,
        format_func=lambda pid:f"{names[pid]} • {pid}",
        key=f"bulk_exception_{project}_{upload_fingerprint}")
    overrides = {}
    if exceptions:
        source_rows = []
        for pid in exceptions:
            row = {"Mã điểm":pid,"Tên địa điểm":names[pid]}
            row.update({sku:float(imported.get((pid,sku), defaults[sku])) for sku in skus})
            source_rows.append(row)
        ui_key = hashlib.sha256(
            repr((project, upload_fingerprint, sorted(exceptions),
                  tuple((k,defaults[k]) for k in skus))).encode()
        ).hexdigest()[:16]
        edited = st.data_editor(
            source_rows, key=f"bulk_exceptions_{ui_key}", hide_index=True,
            use_container_width=True, num_rows="fixed",
            disabled=["Mã điểm", "Tên địa điểm"],
            column_config={
                "Mã điểm":st.column_config.TextColumn("Mã điểm"),
                "Tên địa điểm":st.column_config.TextColumn("Địa điểm"),
                **{sku:st.column_config.NumberColumn(
                    sku, min_value=0.0, step=1.0, format="%.3f")
                   for sku in skus}
            })
        exception_rows = edited.to_dict("records") if hasattr(edited,"to_dict") else edited
        for row in exception_rows:
            pid = row["Mã điểm"]
            for sku in skus:
                overrides[(pid,sku)] = row.get(sku,defaults[sku])
    plan = build_plan([point_by_id[pid] for pid in chosen_ids],
                      skus, defaults, overrides)
    errors, candidates, totals = verify_plan(
        project, plan, inventory, allocations, points)
    st.markdown("#### 4. Đối soát trước khi xác nhận")
    preview_summary = []
    for sku,itm in inventory.items():
        preview_summary.append({
            "SKU":sku,"Tên hàng":itm["name"],"ĐVT":itm["unit"],
            "Tồn khả dụng":float(itm["remaining"]),
            "Dự kiến phân bổ":float(totals.get(sku, 0)),
            "Còn sau phân bổ":float(itm["remaining"]-totals.get(sku,0))
        })
    display_centered_table(preview_summary, hide_index=True, use_container_width=True)
    st.caption(f"Dự kiến ghi {len(candidates)} dòng vào KHO_PHAN_BO; "
               f"phát sinh tương ứng nhiệm vụ VC và KTV. "
               f"Địa điểm được chọn: {len(chosen_ids)}; ngoại lệ: {len(exceptions)}.")
    with st.expander("Xem toàn bộ ma trận địa điểm × SKU",expanded=False):
        display_centered_table([{"Mã điểm":r["Mã điểm"],"Địa điểm":r["Tên địa điểm"],
                       **{sku:float(r["quantities"][sku]) for sku in skus}}
                      for r in plan],hide_index=True,use_container_width=True)
    st.download_button("⬇️ Tải ma trận xem trước (.xlsx)",
        data=bulk_matrix_xlsx(plan, skus),
        file_name=f"PHAN_BO_NHAP_{project}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"bulk_preview_download_{project}")
    if errors:
        for error in errors[:12]:
            st.error(error)
        if len(errors)>12:
            st.error(f"Còn {len(errors)-12} lỗi khác; phải xử lý trước khi ghi.")
        return
    st.warning("Chỉ xác nhận khi bộ SKU, định mức của mỗi điểm và DANH SÁCH "
               "địa điểm đã khớp hồ sơ thầu. Không cho phân bổ trùng SKU/điểm "
               "hoặc vượt hàng thực nhập.")
    if len(candidates)>1600:
        st.info("Đợt lớn có thể tải lâu khi ghi nhiều dòng. "
                "Chỉ bấm xác nhận một lần và đối soát mã đợt.")
    basket_signature=hashlib.sha256(repr((
        project,tuple(sorted(chosen_ids)),
        tuple((r["Mã điểm"],tuple(sorted(r["quantities"].items()))) for r in plan)
    )).encode("utf-8")).hexdigest()[:16]
    accepted=st.checkbox(
        f"Tôi đã kiểm tra {len(chosen_ids)} điểm, định mức SKU, tồn khả dụng "
        "và xác nhận đúng hồ sơ thầu.",
        key=f"bulk_approved_{basket_signature}")
    if not st.button("✅ XÁC NHẬN PHÂN BỔ HÀNG LOẠT",
                     key=f"bulk_confirm_{project}",type="primary",
                     use_container_width=True,disabled=not accepted):
        return
    # A single atomic Google Sheets batch contains the data for all tables.
    # Re-read live without cached data immediately before commit.
    try:
        live_imports = read_sheet_live("NHAP_KHO")
        live_allocations = read_sheet_live("KHO_PHAN_BO")
        live_points = read_sheet_live("DANH_SACH_DIEM")
        live_projects = read_sheet_live("DANH_SACH_DU_AN")
        if project not in {p.get("Mã dự án") for p in live_projects}:
            st.error("Dự án đã thay đổi, không ghi phân bổ.")
            return
        live_inventory, live_errors = inventory_by_sku(
            project,live_imports,live_allocations)
        changed = any(sku not in live_inventory or
            (live_inventory[sku]["name"],live_inventory[sku]["unit"]) !=
            (inventory[sku]["name"],inventory[sku]["unit"])
            for sku in skus)
        if live_errors or changed:
            st.error("Danh mục kho đã thay đổi; vui lòng kiểm tra và lập lại.")
            return
        live_errors, live_candidates, live_totals = verify_plan(
            project,plan,live_inventory,live_allocations,live_points)
        if live_errors:
            for issue in live_errors[:8]:
                st.error(issue)
            st.error("Không ghi dữ liệu. Hãy kiểm tra lại phương án.")
            return
        batch_id="PB-"+uuid.uuid4().hex[:12].upper()
        sheet_names=["DM_CHUAN","KHO_PHAN_BO","VAN_CHUYEN","LAP_DAT",
                     "TIEN_DO","NHAT_KY_HE_THONG"]
        spreadsheet=client()[0]
        ids={name:spreadsheet.worksheet(name).id for name in sheet_names}
        existing_tasks={
            "VAN_CHUYEN":read_sheet_live("VAN_CHUYEN"),
            "LAP_DAT":read_sheet_live("LAP_DAT")
        }
        progress=read_sheet_live("TIEN_DO")
        requests = atomic_requests(
            ids,live_candidates,batch_id,project,live_inventory,
            names,auth()["username"],now(),existing_tasks,progress,
            approved_owners=allocation_owners(
                read_sheet_live("DANG_KY_THANH_VIEN"),
                read_sheet_live("QUAN_LY_DOI"),project))
        spreadsheet.batch_update({"requests":requests})
        read_sheet.clear()
        st.success(f"Đã phân bổ {len(live_candidates)} dòng tại "
                   f"{len({p for p,s,q in live_candidates})} điểm. "
                   f"Mã đợt: {batch_id}. "
                   "Đã tạo nhiệm vụ VC/KTV và tự liên kết đội đã được duyệt; "
                    "điểm chưa có đội duyệt sẽ chờ đội đăng ký.")
    except Exception as exc:
        read_sheet.clear()
        st.error("CHƯA XÁC NHẬN ĐƯỢC KẾT QUẢ GHI ("+type(exc).__name__+
                 "). KHÔNG bấm lại. Kiểm tra KHO_PHAN_BO, VAN_CHUYEN, "
                 "LAP_DAT và NHAT_KY_HE_THONG trước khi thử tiếp.")


def allocate():
    require_role("ADMIN")
    mode=st.radio("Cách phân bổ",[
        "Hàng loạt – định mức chung + ngoại lệ",
        "Theo từng địa điểm (cách cũ)",
        "Điều chỉnh phân bổ đã ghi (ADMIN)"
    ],key="allocation_mode")
    if mode.startswith("Hàng loạt"):
        allocate_bulk()
    elif mode.startswith("Điều chỉnh"):
        adjust_allocated_quantities()
    else:
        allocate_single()


def adjust_allocated_quantities():
    """ADMIN adjusts untouched project allocations, on phone or desktop."""
    require_role("ADMIN")
    st.subheader("✏️ ĐIỀU CHỈNH PHÂN BỔ ĐÃ GHI")
    st.caption("Sửa số lượng từng SKU ở các địa điểm đã phân bổ nhầm. "
               "Chỉ thực hiện trước khi có báo cáo VC/LĐ hoặc biên bản bàn giao. "
               "Số lượng mới được đồng bộ cùng lúc ở Định mức, Phân bổ kho, VC và LĐ.")
    projects = [r["Mã dự án"] for r in read_sheet("DANH_SACH_DU_AN")
                if r.get("Mã dự án")]
    if not projects:
        st.info("Chưa có dự án.")
        return
    project = st.selectbox("1. Dự án", projects,
                           index=projects.index("DA880") if "DA880" in projects else 0,
                           key="adjust_allocation_project")
    allocations = [r for r in read_sheet("KHO_PHAN_BO")
                   if r.get("Mã dự án") == project and r.get("Mã phân bổ")]
    if not allocations:
        st.info("Dự án này chưa có hàng đã phân bổ.")
        return
    batches = sorted({allocation_batch_code(r) for r in allocations
                      if allocation_batch_code(r)}, reverse=True)
    if not batches:
        st.error("Thiếu mã đợt phân bổ; ADMIN cần đối soát dữ liệu.")
        return
    batch = st.selectbox("2. Đợt phân bổ", batches, key=f"adjust_batch_{project}")
    batch_rows = [r for r in allocations if allocation_batch_code(r) == batch]
    skus = sorted({r.get("SKU") for r in batch_rows if r.get("SKU")})
    if not skus:
        st.error("Đợt này thiếu SKU; chưa được điều chỉnh.")
        return
    sku_names={r.get("SKU"):r.get("Tên hàng hóa","") for r in batch_rows}
    sku = st.selectbox("3. Hàng hóa cần điều chỉnh", skus,
                       format_func=lambda code:
                           f"{sku_names.get(code) or code} • SKU {code}",
                       key=f"adjust_sku_{project}_{batch}")
    matches = [r for r in batch_rows if r.get("SKU") == sku]
    by_alloc = {r["Mã phân bổ"]: r for r in matches}
    if len(by_alloc) != len(matches):
        st.error("Có mã phân bổ trùng; cần đối soát trước khi sửa.")
        return
    point_names = {r.get("Mã điểm"): r.get("Tên xã/phường/đơn vị")
                   for r in read_sheet("DANH_SACH_DIEM")
                   if r.get("Mã dự án") == project}
    st.caption("Đối chiếu toàn bộ địa điểm trong đợt/SKU này; số lượng bên dưới là số đang lưu, chưa phải định mức đúng.")
    display_centered_table([
        {"Tên đơn vị": point_names.get(r.get("Mã điểm")) or "(Chưa có trong danh sách địa điểm)",
         "Mã điểm": r.get("Mã điểm", ""),
         "Tên hàng hóa": sku_names.get(sku,""),
         "SKU": sku,
         "SL đang phân bổ": r.get("SL phân bổ", ""),
         "ĐVT": r.get("ĐVT", ""),
         "Mã phân bổ": r.get("Mã phân bổ", "")}
        for r in sorted(matches, key=lambda r:
            (str(point_names.get(r.get("Mã điểm")) or ""), str(r.get("Mã điểm") or "")))
    ], hide_index=True, use_container_width=True)

    chosen = st.multiselect(
        "4. Địa điểm cần sửa (chọn nhiều điểm nếu cùng số lượng mới)",
        list(by_alloc),
        format_func=lambda aid: (
            f"{point_names.get(by_alloc[aid].get('Mã điểm'), by_alloc[aid].get('Mã điểm'))}"
            f" • {by_alloc[aid].get('Mã điểm')}"
            f" • {by_alloc[aid].get('SL phân bổ')} {by_alloc[aid].get('ĐVT', '')}"),
        key=f"adjust_ids_{project}_{batch}_{sku}")
    if not chosen:
        st.info("Chọn địa điểm để xem số cũ và sửa số đúng.")
        return
    selection_fingerprint = hashlib.sha256(
        repr((project, batch, sku, sorted(
            (aid, by_alloc[aid].get("SL phân bổ")) for aid in chosen))).encode()
    ).hexdigest()[:12]
    new_qty = st.number_input(
        "5. ADMIN NHẬP SỐ LƯỢNG ĐÚNG CHO MỖI ĐỊA ĐIỂM ĐÃ CHỌN",
        min_value=0.001, value=None, step=1.0, format="%.3f",
        placeholder="ADMIN nhập số lượng đúng",
        key=f"adjust_qty_{selection_fingerprint}")
    if new_qty is None:
        st.info("ADMIN nhập số lượng đúng để lập bảng đối chiếu trước khi xác nhận.")
        return
    changes = [(aid, by_alloc[aid].get("SL phân bổ"), new_qty) for aid in chosen]
    preview = []
    for aid, old, updated in changes:
        try:
            previous = float(old)
            amount = previous - updated
        except (ValueError, TypeError):
            st.error("Dữ liệu số lượng cũ không hợp lệ: " + str(aid))
            return
        preview.append({"Địa điểm": point_names.get(by_alloc[aid].get("Mã điểm"),
                                                    by_alloc[aid].get("Mã điểm")),
                        "SKU": sku, "SL cũ": previous,
                        "SL mới": updated, "Trả lại kho phân bổ": amount})
    display_centered_table(preview, hide_index=True, use_container_width=True)
    if all(row["SL cũ"] == row["SL mới"] for row in preview):
        st.warning("Số mới chưa khác số cũ.")
        return
    st.warning("Chỉ sửa phân bổ CHƯA PHÁT SINH thực hiện. Nếu đã giao/lắp hoặc "
               "đã có biên bản, hệ thống từ chối để tránh lệch tồn kho.")
    signature = hashlib.sha256(repr((project, batch, sku, changes)).encode()).hexdigest()[:12]
    accepted = st.checkbox("Tôi đã đối chiếu số lượng thực tế cần phân bổ cho "
                           "đúng dự án, SKU và từng địa điểm.",
                           key=f"adjust_confirm_{signature}")
    if not st.button("✅ XÁC NHẬN ĐIỀU CHỈNH PHÂN BỔ",
                     type="primary", use_container_width=True,
                     key=f"adjust_commit_{signature}", disabled=not accepted):
        return
    try:
        required = ("NHAP_KHO", "KHO_PHAN_BO", "DM_CHUAN", "VAN_CHUYEN",
                    "LAP_DAT", "BAO_CAO_TRIEN_KHAI", "BAN_GIAO_TAP_KET")
        snapshots = {name: read_sheet_live(name) for name in required}
        spreadsheet = client()[0]
        wanted = {"DM_CHUAN", "KHO_PHAN_BO", "VAN_CHUYEN",
                  "LAP_DAT", "NHAT_KY_HE_THONG"}
        sheet_ids = {ws.title: ws.id for ws in spreadsheet.worksheets()
                     if ws.title in wanted}
        requests, delta = prepare_quantity_corrections(
            project, changes, snapshots, sheet_ids, auth()["username"], now())
        spreadsheet.batch_update({"requests": requests})
        read_sheet.clear()
        returned = sum(float(-qty) for qty in delta.values() if qty < 0)
        st.success(f"Đã điều chỉnh {len(changes)} địa điểm/SKU trong một lần ghi. "
                   f"Số lượng trả lại để phân bổ: {returned:g}. "
                   "Các bảng Định mức, Kho, VC, LĐ và Nhật ký đã đồng bộ.")
    except ValueError as exc:
        read_sheet.clear()
        st.error("CHƯA GHI ĐIỀU CHỈNH: " + str(exc))
    except Exception as exc:
        read_sheet.clear()
        st.error("CHƯA XÁC MINH ĐƯỢC KẾT QUẢ GHI (" + type(exc).__name__ +
                 "). KHÔNG BẤM LẠI ngay; ADMIN cần đối chiếu "
                 "DM_CHUAN, KHO_PHAN_BO, VAN_CHUYEN, LAP_DAT và Nhật ký.")


def allocate_single():
    """ADMIN single-location allocation, using the same atomic path as bulk."""
    require_role("ADMIN")
    st.subheader("📦 Phân bổ SLHH theo từng điểm")
    st.caption("Chỉ phân bổ hàng; không ghi xuất kho hoặc hoàn thành lắp đặt.")
    projects=[r["Mã dự án"] for r in read_sheet("DANH_SACH_DU_AN")
              if r.get("Mã dự án")]
    if not projects:
        st.info("Cần tạo dự án trước.")
        return
    project=st.selectbox("Dự án",projects)
    points=[x for x in read_sheet("DANH_SACH_DIEM")
            if x.get("Mã dự án")==project and x.get("Mã điểm")
            and x.get("Trạng thái")=="Hoạt động"]
    if not points:
        st.info("Cần có điểm đang hoạt động thuộc dự án này.")
        return
    if len(points)!=len({x["Mã điểm"] for x in points}):
        st.error("Danh sách địa điểm trùng mã; ADMIN cần đối soát trước.")
        return
    pmap={f"{x['Tên xã/phường/đơn vị']} • {x['Mã điểm']}":x for x in points}
    point=pmap[st.selectbox("Điểm cần phân bổ",list(pmap))]
    inventory,issues=inventory_by_sku(
        project,read_sheet("NHAP_KHO"),read_sheet("KHO_PHAN_BO"))
    if issues:
        for issue in issues:
            st.error(issue)
        st.warning("Cần đối soát hàng nhập và phân bổ trước khi tiếp tục.")
        return
    if not inventory:
        st.info("Chưa có hàng nhập kho của dự án.")
        return
    with st.form("allocate_batch"):
        quantities={}
        for sku,item in inventory.items():
            remain=max(0,item["remaining"])
            st.caption(f"{sku} • {item['name']} • Hàng còn chờ phân bổ: "
                       f"{remain:g} {item['unit']}")
            quantities[sku]=st.number_input(
                f"SL phân bổ {sku}",min_value=0.0,
                max_value=float(remain),step=1.0,
                key=f"alloc_{project}_{point['Mã điểm']}_{sku}")
        submitted=st.form_submit_button("PHÂN BỔ",use_container_width=True)
    if not submitted:
        return

    chosen={sku:qty for sku,qty in quantities.items() if qty>0}
    if not chosen:
        st.warning("Chưa nhập SLHH.")
        return
    proposed=[{"Mã điểm":point["Mã điểm"],"quantities":chosen}]
    errors,candidates,_=verify_plan(
        project,proposed,inventory,read_sheet("KHO_PHAN_BO"),points)
    if errors:
        for issue in errors:
            st.error(issue)
        return
    batch_id="PB-"+uuid.uuid4().hex[:12].upper()
    try:
        # Re-read LIVE sheets before any write; same validation as bulk mode.
        spreadsheet=client()[0]
        live_projects=read_sheet_live("DANH_SACH_DU_AN")
        if project not in {r.get("Mã dự án") for r in live_projects}:
            raise ValueError("Dự án không còn tồn tại.")
        live_points=read_sheet_live("DANH_SACH_DIEM")
        live_allocations=read_sheet_live("KHO_PHAN_BO")
        live_inventory,inventory_issues=inventory_by_sku(
            project,read_sheet_live("NHAP_KHO"),live_allocations)
        if inventory_issues:
            raise ValueError("Tồn hàng cần đối soát: "+
                             "; ".join(inventory_issues[:3]))
        live_errors,live_candidates,_=verify_plan(
            project,proposed,live_inventory,live_allocations,live_points)
        if live_errors:
            raise ValueError("; ".join(live_errors[:3]))
        if live_candidates!=candidates:
            raise ValueError("Phân bổ đã thay đổi trong lúc xác nhận.")
        point_live=next((r for r in live_points
                         if r.get("Mã dự án")==project
                         and r.get("Mã điểm")==point["Mã điểm"]),None)
        if point_live is None:
            raise ValueError("Không còn tìm thấy địa điểm trong dự án.")
        tables=("DM_CHUAN","KHO_PHAN_BO","VAN_CHUYEN","LAP_DAT",
                "TIEN_DO","NHAT_KY_HE_THONG")
        ids={name:spreadsheet.worksheet(name).id for name in tables}
        existing_tasks={"VAN_CHUYEN":read_sheet_live("VAN_CHUYEN"),
                        "LAP_DAT":read_sheet_live("LAP_DAT")}
        reqs=atomic_requests(
            ids,live_candidates,batch_id,project,live_inventory,
            {point["Mã điểm"]:point_live["Tên xã/phường/đơn vị"]},
            auth()["username"],now(),existing_tasks,
            read_sheet_live("TIEN_DO"),
            approved_owners=allocation_owners(
                read_sheet_live("DANG_KY_THANH_VIEN"),
                read_sheet_live("QUAN_LY_DOI"),project))
        # One all-or-nothing Sheets transaction: NO sequential appends.
        spreadsheet.batch_update({"requests":reqs})
        read_sheet.clear()
        st.success(f"Đã phân bổ đồng bộ {len(live_candidates)} SKU tại "
                   f"{point_live['Tên xã/phường/đơn vị']}. "
                   "Nhiệm vụ VC/KTV đã tự liên kết đội có địa điểm được duyệt.")
    except ValueError as exc:
        read_sheet.clear()
        st.error("Chưa phân bổ: "+str(exc))
    except Exception as exc:
        read_sheet.clear()
        st.error("CHƯA XÁC NHẬN ĐƯỢC KẾT QUẢ GHI ("+type(exc).__name__+
                 "). KHÔNG bấm lại. ADMIN cần đối chiếu các sheet liên quan.")


def provision_team_password():
    """ADMIN-only provisioning. Password is shown ONCE and never stored clear.

    Existing signed remembered logins become invalid after a password reset.
    The sheet is an ADMIN/SA-only credential vault; do not share edit access.
    """
    a=require_role("ADMIN")
    with st.expander("🔐 CẤP / ĐẶT LẠI MẬT KHẨU – ĐỘI ĐÃ TẠO",
                     expanded=False):
        st.caption("Danh sách này đọc từ QUAN_LY_DOI (đội đã có trong hệ thống), "
                   "KHÔNG đọc từ DANG_KY_THANH_VIEN. "
                   "Xóa hồ sơ đăng ký không xóa các đội này. "
                   "Hồ sơ mới có mật khẩu tự chọn sẽ kích hoạt khi ADMIN duyệt; "
                   "đội cũ hoặc quên mật khẩu có thể cấp/đặt lại tại đây.")
        teams=[t for t in read_sheet("QUAN_LY_DOI")
               if t.get("Trạng thái")=="Hoạt động"
               and t.get("Loại nhiệm vụ") in ("VC","KTV","VC+KTV")
               and t.get("Mã dự án") and t.get("Mã đội")]
        if not teams:
            st.info("Chưa có đội trưởng VC/KTV đã duyệt. "
                    "Cần duyệt hồ sơ trước khi cấp mật khẩu.")
            return
        team=st.selectbox(
            "Chọn đội trưởng đang có trong QUAN_LY_DOI", teams,
            format_func=lambda t:(t.get("Tên đội/đội trưởng","")+
                                  " • "+t.get("Mã dự án","")+
                                  " • "+t.get("Loại nhiệm vụ","")),
            key="credential_team_selection")
        if sum(normalized_name(t.get("Tên đội/đội trưởng"))==
               normalized_name(team.get("Tên đội/đội trưởng"))
               for t in teams)>1:
            st.error("Có nhiều đội trưởng trùng họ tên. "
                     "Hiện đăng nhập bằng tên không phân biệt được họ; "
                     "ADMIN cần làm rõ định danh trước khi cấp.")
            return
        records=read_sheet("TAI_KHOAN_DOI")
        current,existing=visible_credential_state(
            team,records,st.secrets.get("accounts",{}))
        if current=="conflict":
            st.error("Có nhiều tài khoản trùng mã đội. "
                     "Cần xử lý hồ sơ tài khoản trước khi cấp.")
            return
        if current=="legacy":
            st.info("Đội này đã có tài khoản trong Streamlit Secrets. "
                    "Để tránh hai tài khoản cùng quyền, B chưa cho cấp thêm "
                    "ở đây; ADMIN cần đối soát tài khoản cũ trước.")
            return
        if current=="active":
            st.info("Đội trưởng đã được cấp tài khoản. "
                    "Nếu quên mật khẩu, có thể ĐẶT LẠI; "
                    "mật khẩu cũ sẽ mất hiệu lực.")
        else:
            st.info("Đội trưởng chưa có tài khoản. Có thể cấp mới.")
        action="ĐẶT LẠI MẬT KHẨU" if current=="active" else "CẤP MẬT KHẨU"
        method=st.radio(
            "Cách chọn mật khẩu",
            ["Đội trưởng đề nghị / ADMIN nhập hộ","Hệ thống tự tạo mật khẩu"],
            horizontal=False,key="credential_password_method")
        # Submit both fields as one Streamlit form to avoid stale widget values.
        with st.form("credential_provision_form"):
            if method.startswith("Đội trưởng"):
                st.caption("Chọn mật khẩu dễ nhớ, từ 6 ký tự; "
                           "có chữ và số. Gửi riêng cho người nhận.")
                custom_password=st.text_input(
                    "Mật khẩu theo đề nghị *",type="password",
                    max_chars=128,key="credential_custom_password")
                confirm_custom=st.text_input(
                    "Nhập lại mật khẩu *",type="password",
                    max_chars=128,key="credential_custom_confirm")
            approve=st.checkbox(
                "Tôi xác nhận đúng người nhận và sẽ chỉ gửi mật khẩu "
                "qua tin nhắn riêng, không đăng vào nhóm.",
                key="credential_confirm_private")
            submitted=st.form_submit_button(
                "🔐 "+action,type="primary",use_container_width=True)
        if not submitted:
            return
        if not approve:
            st.error("Vui lòng xác nhận đúng người nhận mật khẩu.")
            return
        if method.startswith("Đội trưởng"):
            if custom_password!=confirm_custom:
                st.error("Hai lần nhập mật khẩu chưa giống nhau. "
                         "Chưa cấp/đặt lại mật khẩu.")
                return
            message=validate_custom_password(
                custom_password,team.get("Tên đội/đội trưởng",""),
                team.get("SĐT đội trưởng",""))
            if message:
                st.error(message)
                return
        try:
            # Strict live recheck before writing: sheet and team can change
            # while the ADMIN form is open.
            read_sheet.clear()
            teams_live=read_sheet("QUAN_LY_DOI")
            live=[t for t in teams_live
                  if t.get("Mã đội")==team["Mã đội"]
                  and t.get("Mã dự án")==team["Mã dự án"]
                  and t.get("Trạng thái")=="Hoạt động"
                  and t.get("Loại nhiệm vụ") in ("VC","KTV","VC+KTV")]
            if len(live)!=1 or live[0].get("Tên đội/đội trưởng")!=team.get("Tên đội/đội trưởng"):
                fail("Thông tin đội trưởng đã thay đổi. Hãy tải lại.")
            if sum(normalized_name(t.get("Tên đội/đội trưởng"))==
                   normalized_name(team.get("Tên đội/đội trưởng"))
                   for t in teams_live if t.get("Trạng thái")=="Hoạt động")>1:
                fail("Tên đội trưởng trùng; phải đối soát trước khi cấp.")
            fresh=read_sheet("TAI_KHOAN_DOI")
            kind,row=visible_credential_state(
                live[0],fresh,st.secrets.get("accounts",{}))
            if kind!=current or kind in ("conflict","legacy"):
                fail("Tình trạng tài khoản thay đổi; hãy mở lại.")
            password=(custom_password if method.startswith("Đội trưởng")
                      else new_password())
            salt,digest=create_hash(password)
            admin=str(a["username"])
            when=now()
            if current=="active":
                # Replace only salted secret, status, time and audit owner.
                # Never duplicate active accounts for a team.
                idx=row.get("_row")
                if not isinstance(idx,int) or idx<2:
                    fail("Không tìm thấy dòng tài khoản để đặt lại.")
                sheet("TAI_KHOAN_DOI").update(
                    range_name=f"E{idx}:J{idx}",
                    values=[[salt,digest,"Hoạt động",when,admin,
                             "ADMIN đặt lại mật khẩu trên DỰ ÁN"]],
                    value_input_option="RAW")
                username=row["Mã tài khoản"]
            else:
                usernames={r.get("Mã tài khoản") for r in fresh}
                usernames.update(st.secrets.get("accounts",{}).keys())
                username=next_account_id(usernames)
                append("TAI_KHOAN_DOI", [
                    username,team["Mã đội"],team["Mã dự án"],
                    "DOI_TRUONG",salt,digest,"Hoạt động",when,
                    admin,"ADMIN cấp mật khẩu trên DỰ ÁN"])
            read_sheet.clear()
            verified=effective_accounts().get(username)
            if not verified or verified.get("hash")!=digest or (
                    verified.get("team_id")!=team["Mã đội"]):
                fail("Đã ghi nhưng chưa xác nhận tài khoản hợp lệ. "
                     "Không cấp mật khẩu; hãy đối soát trước khi thử lại.")
            # No raw password in Sheets, GitHub, audit logs, or cookies.
            audit(admin,"CẤP MẬT KHẨU ĐỘI TRƯỞNG","TAI_KHOAN_DOI",
                  username,"OK",team["Mã dự án"]+
                  " / "+("Đặt lại" if current=="active" else "Cấp mới"))
            st.success("Đã "+("đặt lại" if current=="active" else "cấp")+
                       " mật khẩu cho **"+team["Tên đội/đội trưởng"]+"**.")
            st.warning("Mật khẩu chỉ được hiển thị cho ADMIN tại lần cấp này. "
                       "Hãy gửi riêng cho đội trưởng; không đăng nhóm.")
            st.code(password,language=None)
            st.caption("Đội trưởng mở app → ĐỘI TRƯỞNG – BÁO CÁO "
                       "→ nhập HỌ TÊN đã đăng ký + mật khẩu này. "
                       "Có thể tích Ghi nhớ đăng nhập 30 ngày trên máy cá nhân.")
        except Exception as exc:
            read_sheet.clear()
            st.error("Chưa xác nhận được việc cấp mật khẩu ("+
                     type(exc).__name__+"). Không gửi bất kỳ mật khẩu "
                     "nào nếu chưa thấy thông báo cấp thành công. "
                     "Cần kiểm tra bảng tài khoản trước khi làm lại.")



def supplemental_site_assignment_admin():
    """ADMIN adds previously unowned project sites to an approved VC/KTV team."""
    require_role("ADMIN")
    with st.expander("➕ GIAO BỔ SUNG ĐIỂM CHƯA CÓ ĐỘI (VC / LẮP ĐẶT)",
                     expanded=False):
        notice = st.session_state.pop("supplemental_site_notice", None)
        if notice:
            st.success(notice)
        st.caption("Chọn đội đang hoạt động và điểm chưa cấp cho đội nào ở "
                   "đúng công đoạn. Giữ nguyên các điểm đã giao, số lượng "
                   "kho và mọi báo cáo/biên bản. Có thể làm trên PC hoặc điện thoại.")
        all_teams = read_sheet("QUAN_LY_DOI")
        teams = [r for r in all_teams
                 if r.get("Mã đội") and r.get("Mã dự án")
                 and r.get("Trạng thái") == "Hoạt động" and allowed_modes(r)]
        if not teams:
            st.info("Chưa có đội VC/LĐ đang hoạt động.")
            return
        project = st.selectbox("1. Dự án", sorted({r["Mã dự án"] for r in teams}),
                               key="supplemental_project")
        team = st.selectbox(
            "2. Đội nhận thêm điểm",
            [r for r in teams if r["Mã dự án"] == project],
            format_func=lambda r: r.get("Tên đội/đội trưởng", r["Mã đội"])
            + " • " + r.get("Loại nhiệm vụ", ""),
            key=f"supplemental_team_{project}")
        mode_labels = {"VC": "Vận chuyển", "LD": "Lắp đặt",
                       "VC+LD": "Cả VC và lắp đặt"}
        mode = st.radio("3. Công đoạn bổ sung", allowed_modes(team),
                        format_func=lambda m: mode_labels[m],
                        key=f"supplemental_mode_{team['Mã đội']}")
        tabs = ROLE_TABS[mode]
        locations = read_sheet("DANH_SACH_DIEM")
        registrations = read_sheet("DANG_KY_THANH_VIEN")
        tasksets = {tab: read_sheet(tab) for tab in tabs}
        hubs = read_sheet("BAN_GIAO_TAP_KET")
        try:
            active_registration(registrations, project, team["Mã đội"])
            owners = owner_index(registrations, all_teams, project)
            catalog = project_site_catalog(team, mode, locations, tasksets)
            # Check EVERY candidate with the same pure planner used at commit.
            ready = {}
            for pid, (loc, status, _) in catalog.items():
                if any(
                        (tab, pid) in owners for tab in tabs):
                    continue
                try:
                    supplemental_site_plan(
                        team, mode, (pid,), locations, registrations,
                        all_teams, tasksets, hubs)
                    ready[pid] = loc
                except ValueError:
                    continue
        except ValueError as exc:
            st.warning("Cần đối soát quyền đội trước khi bổ sung: " + str(exc))
            return
        st.caption(f"**{len(ready)} điểm** chưa cấp quyền cho đội nào "
                   "ở công đoạn này. Gồm cả điểm chưa phân bổ SKU.")
        if not ready:
            st.info("Không có điểm trống hợp lệ để bổ sung. Điểm đã thuộc "
                    "đội khác phải qua mục Thu hồi / Giao lại đội khác.")
            return
        chosen = st.multiselect(
            "4. Tìm tên đơn vị / điểm cần bổ sung",
            sorted(ready, key=lambda pid:
                   ready[pid].get("Tên xã/phường/đơn vị", pid).casefold()),
            format_func=lambda pid: (
                ready[pid].get("Tên xã/phường/đơn vị") or pid) + " • " + pid,
            placeholder="Gõ tên đơn vị hoặc mã điểm...",
            key=f"supplemental_sites_{project}_{team['Mã đội']}_{mode}")
        if not chosen:
            st.info("Chọn một hoặc nhiều điểm để xem đủ SKU và SL trước khi giao.")
            return
        try:
            plan = supplemental_site_plan(
                team, mode, chosen, locations, registrations,
                all_teams, tasksets, hubs)
        except ValueError as exc:
            st.warning(str(exc))
            return
        stock = [r for r in read_sheet("KHO_PHAN_BO")
                 if r.get("Mã dự án") == project and r.get("Mã điểm") in chosen]
        stock_by_key = {(r.get("Mã điểm"), r.get("SKU")): r for r in stock}
        st.markdown("#### 5. Đối chiếu SKU và số lượng từng địa điểm")
        for pid in chosen:
            loc = ready[pid]
            quantities = {}
            for tab, rows in plan["tasks"].items():
                for row in rows:
                    if row.get("Mã điểm") != pid:
                        continue
                    sku = row.get("SKU")
                    entry = quantities.setdefault(sku, {"SKU": sku})
                    entry["SL VC" if tab == "VAN_CHUYEN" else "SL LĐ"] = (
                        row.get("SL giao") if tab == "VAN_CHUYEN"
                        else row.get("SL lắp"))
            preview = []
            for sku, values in sorted(quantities.items()):
                source = stock_by_key.get((pid, sku), {})
                display = {"SKU": sku,
                           "Thiết bị": source.get("Tên hàng hóa", ""),
                           "ĐVT": source.get("ĐVT", "")}
                if "SL VC" in values:
                    display["SL giao"] = values["SL VC"]
                if "SL LĐ" in values:
                    display["SL lắp"] = values["SL LĐ"]
                preview.append(display)
            with st.container(border=True):
                st.markdown("**📍 " +
                            (loc.get("Tên xã/phường/đơn vị") or pid) + "**")
                st.caption(f"Mã điểm: {pid} • {len(preview)} SKU")
                if preview:
                    display_centered_table(preview, hide_index=True, use_container_width=True)
                else:
                    st.info("Điểm chưa phân bổ SKU. Sau khi ADMIN phân bổ "
                            "hàng, nhiệm vụ sẽ liên kết theo quyền vừa cấp.")
        count = sum(len(rows) for rows in plan["tasks"].values())
        st.caption(f"Tổng cộng: {len(chosen)} điểm, {count} nhiệm vụ SKU. "
                   "Bổ sung cả quyền địa điểm đã duyệt và nhiệm vụ trong cùng "
                   "một lần ghi; không thay đổi phân bổ kho.")
        signature = hashlib.sha256(
            repr(supplemental_site_signature(plan)).encode("utf-8")
        ).hexdigest()[:16]
        checked = st.checkbox(
            "ADMIN xác nhận các điểm trên chưa thuộc đội nào và đúng hợp đồng.",
            key=f"supplemental_ack_{signature}")
        if st.button("✅ XÁC NHẬN GIAO BỔ SUNG ĐIỂM",
                     disabled=not checked, type="primary",
                     use_container_width=True,
                     key=f"supplemental_commit_{signature}"):
            try:
                read_sheet.clear()
                fresh_teams = read_sheet_live("QUAN_LY_DOI")
                matches = [r for r in fresh_teams if
                           r.get("Mã đội") == team["Mã đội"]
                           and r.get("Mã dự án") == project]
                if len(matches) != 1:
                    fail("Đội nhận điểm đã thay đổi; chưa ghi.")
                fresh = supplemental_site_plan(
                    matches[0], mode, chosen,
                    read_sheet_live("DANH_SACH_DIEM"),
                    read_sheet_live("DANG_KY_THANH_VIEN"), fresh_teams,
                    {tab: read_sheet_live(tab) for tab in tabs},
                    read_sheet_live("BAN_GIAO_TAP_KET"))
                if (supplemental_site_signature(fresh) !=
                        supplemental_site_signature(plan)):
                    fail("Quyền hoặc nhiệm vụ đã đổi; chưa ghi. Hãy tải lại.")
                ids = {tab: sheet(tab).id for tab in tabs}
                ids["DANG_KY_THANH_VIEN"] = sheet("DANG_KY_THANH_VIEN").id
                requests = supplemental_site_requests(fresh, ids)
                client()[0].batch_update({"requests": requests})
                read_sheet.clear()
                st.session_state["supplemental_site_notice"] = (
                    f"Đã bổ sung {len(chosen)} điểm và {count} nhiệm vụ "
                    f"cho đội {team.get('Tên đội/đội trưởng', team['Mã đội'])}. "
                    "Quyền địa điểm và nhiệm vụ đã ghi cùng giao dịch; "
                    "KHO_PHAN_BO giữ nguyên.")
                st.rerun()
            except Exception as exc:
                read_sheet.clear()
                st.error("CHƯA XÁC NHẬN ĐƯỢC LẦN GIAO BỔ SUNG (" +
                         type(exc).__name__ + "). KHÔNG bấm lại ngay. "
                         "ADMIN đối chiếu DANG_KY_THANH_VIEN, "
                         "VAN_CHUYEN và LAP_DAT trước khi thử lại.")


def team_reassignment_admin():
    """Explicit ADMIN operation; never alter KHO_PHAN_BO or completed evidence."""
    require_role("ADMIN")
    with st.expander("♻️ THU HỒI PHÂN CÔNG / GIAO LẠI ĐỘI KHÁC",
                     expanded=False):
        notice=st.session_state.pop("team_change_notice",None)
        if notice:
            st.success(notice)
        teams=[r for r in read_sheet("QUAN_LY_DOI")
               if r.get("Mã đội") and r.get("Mã dự án")
               and r.get("Trạng thái")=="Hoạt động"
               and allowed_modes(r)]
        if not teams:
            st.info("Chưa có đội VC/KTV đang hoạt động.")
            return
        projects=sorted(set(r["Mã dự án"] for r in teams))
        project=st.selectbox("1. Dự án",projects,key="team_change_project")
        candidates=[r for r in teams if r["Mã dự án"]==project]
        source=st.selectbox("2. Đội đang phụ trách",
                            candidates,
                            format_func=lambda t:t.get("Tên đội/đội trưởng",t["Mã đội"])
                            +" • "+t.get("Loại nhiệm vụ",""),
                            key=f"team_change_source_{project}")
        actions=("Thu hồi – trả nhiệm vụ về chờ cấp lại",
                 "Giao lại trực tiếp cho đội khác",
                 "Đội nghỉ/ngừng hoạt động – khóa tài khoản và nhả việc chưa làm")
        action=st.radio("3. Hình thức xử lý",actions,
                         key=f"team_change_action_{project}")
        suspend=action.startswith("Đội nghỉ")
        transfer=action.startswith("Giao lại")
        tasks={"VAN_CHUYEN":read_sheet("VAN_CHUYEN"),
               "LAP_DAT":read_sheet("LAP_DAT")}
        hubs=read_sheet("BAN_GIAO_TAP_KET")
        registrations=read_sheet("DANG_KY_THANH_VIEN")
        try:
            source_registration=active_registration(
                registrations,project,source["Mã đội"])
        except ValueError as exc:
            st.warning("Đội chưa có quyền địa điểm đã duyệt: "+str(exc))
            return
        chosen=[]
        target=None
        mode=""
        if suspend:
            st.warning("Đội sẽ ngừng đăng nhập. Mọi nhiệm vụ CHƯA LÀM, "
                       "chưa có chứng từ sẽ được nhả; việc đang làm dở "
                       "hoặc có bàn giao được giữ lại để đối soát.")
        else:
            labels={"VC":"Vận chuyển", "LD":"Lắp đặt",
                    "VC+LD":"Vận chuyển và lắp đặt"}
            modes=allowed_modes(source)
            mode=st.radio("4. Loại công việc cần chuyển",modes,
                          format_func=lambda value:labels[value],
                          key=f"team_change_mode_{project}_{source['Mã đội']}")
            kinds=ROLE_TABS[mode]
            rights=[set(site_ids(source_registration.get(
                RIGHT_FIELDS[tab][0]))) for tab in kinds]
            sites=sorted(set.intersection(*rights) if rights else set())
            if not sites:
                st.info("Đội không có nhiệm vụ thuộc loại đang chọn.")
                return
            locs={(r.get("Mã dự án"),r.get("Mã điểm")):
                  r.get("Tên xã/phường/đơn vị","")
                  for r in read_sheet("DANH_SACH_DIEM")}
            chosen=st.multiselect(
                "5. Tìm và chọn các điểm cần thu hồi / giao lại",
                sites,
                format_func=lambda pid:
                    f"{locs.get((project,pid),pid)} • {pid}",
                placeholder="Gõ tên điểm hoặc mã điểm để tìm...",
                key=f"team_change_sites_{project}_{source['Mã đội']}_{mode}")
            if transfer:
                destinations=[t for t in candidates
                              if t["Mã đội"]!=source["Mã đội"]
                              and mode in allowed_modes(t)]
                if not destinations:
                    st.warning("Chưa có đội khác cùng dự án, đúng chuyên môn "
                               "và đang hoạt động để nhận việc.")
                    return
                target=st.selectbox(
                    "6. Đội sẽ nhận việc",
                    destinations,
                    format_func=lambda t:t.get("Tên đội/đội trưởng",t["Mã đội"])
                    +" • "+t.get("Loại nhiệm vụ",""),
                    key=f"team_change_target_{project}_{source['Mã đội']}_{mode}")
            if not chosen:
                st.info("Chọn một hoặc nhiều điểm để xem nhiệm vụ được phép chuyển.")
                return
        try:
            planned=team_change_plan(
                source,chosen,mode,tasks,hubs,
                target_team=target,suspend=suspend,
                registrations=registrations,
                all_teams=read_sheet("QUAN_LY_DOI"))
        except ValueError as exc:
            st.warning(str(exc))
            return
        c1,c2=st.columns(2)
        c1.metric("Nhiệm vụ có thể chuyển",len(planned["changes"]))
        c2.metric("Nhiệm vụ giữ đối soát",len(planned["protected"]))
        st.caption(f"Điểm có thể nhả/chuyển: {planned['released_sites']}. "
                   "Không đổi SL phân bổ cho điểm; không xóa báo cáo hay ảnh.")
        if planned["changes"]:
            with st.expander("Xem các nhiệm vụ sẽ thay đổi"):
                display_centered_table([{"Công đoạn":r["tab"],"Mã điểm":r["point"],
                              "SKU":r["sku"],"Xử lý":
                              ("Giao cho "+planned["target_id"]
                               if planned["target_id"] else "Chờ phân công lại")}
                             for r in planned["changes"]],
                             hide_index=True,use_container_width=True)
        if planned["protected"]:
            st.warning("Các dòng đã làm/có chứng từ không bị chuyển tự động.")
            with st.expander("Xem nhiệm vụ phải bàn giao / đối soát"):
                display_centered_table([{"Công đoạn":r["tab"],"Mã điểm":r["point"],
                              "SKU":r["sku"],"Lý do":r["reason"]}
                             for r in planned["protected"]],
                             hide_index=True,use_container_width=True)
        digest=hashlib.sha256(repr(change_signature(planned)).encode()).hexdigest()[:16]
        ack=st.checkbox(
            "Tôi đã kiểm tra đúng đội, địa điểm và các nhiệm vụ sẽ thay đổi. "
            "Tôi đồng ý thực hiện trên DỰ ÁN.",
            key=f"team_change_confirm_{digest}")
        label=("✅ XÁC NHẬN NGỪNG ĐỘI / THU HỒI VIỆC" if suspend else
               "✅ XÁC NHẬN GIAO LẠI ĐỘI KHÁC" if transfer else
               "✅ XÁC NHẬN TRẢ VỀ CHỜ PHÂN CÔNG")
        if st.button(label,type="primary",use_container_width=True,
                     disabled=not ack,key=f"team_change_commit_{digest}"):
            try:
                read_sheet.clear()
                live_teams=read_sheet_live("QUAN_LY_DOI")
                matches=[r for r in live_teams
                         if r.get("Mã đội")==source["Mã đội"]
                         and r.get("Mã dự án")==project]
                if len(matches)!=1:
                    fail("Không xác định được đội nguồn duy nhất.")
                live_source=matches[0]
                live_target=None
                if transfer:
                    matches=[r for r in live_teams
                             if r.get("Mã đội")==target["Mã đội"]
                             and r.get("Mã dự án")==project]
                    if len(matches)!=1:
                        fail("Không xác định được đội nhận duy nhất.")
                    live_target=matches[0]
                fresh_tasks={"VAN_CHUYEN":read_sheet_live("VAN_CHUYEN"),
                             "LAP_DAT":read_sheet_live("LAP_DAT")}
                live_plan=team_change_plan(
                    live_source,chosen,mode,fresh_tasks,
                    read_sheet_live("BAN_GIAO_TAP_KET"),
                    target_team=live_target,suspend=suspend,
                    registrations=read_sheet_live("DANG_KY_THANH_VIEN"),
                    all_teams=live_teams)
                if change_signature(live_plan)!=change_signature(planned):
                    fail("Nhiệm vụ hoặc bằng chứng đã thay đổi. Tải lại để đối soát.")
                account_rows=[]
                link_rows=[]
                if suspend:
                    accounts=[r for r in read_sheet_live("TAI_KHOAN_DOI")
                              if r.get("Mã đội")==source["Mã đội"]
                              and r.get("Mã dự án")==project
                              and r.get("Trạng thái")=="Hoạt động"]
                    if len(accounts)>1:
                        fail("Đội có nhiều tài khoản hoạt động; cần đối soát.")
                    account_rows=[r.get("_row") for r in accounts]
                    links=[r for r in read_sheet_live("LIEN_KET_DANG_KY_DOI")
                           if r.get("Mã đội")==source["Mã đội"]
                           and r.get("Mã dự án")==project
                           and r.get("Trạng thái liên kết")==LINK_ACTIVE]
                    if len(links)>1:
                        fail("Đội có nhiều hồ sơ liên kết; cần đối soát.")
                    link_rows=[r.get("_row") for r in links]
                sheetnames=["VAN_CHUYEN","LAP_DAT","NHAT_KY_HE_THONG",
                            "DANG_KY_THANH_VIEN"]
                if suspend:
                    sheetnames+=["QUAN_LY_DOI","TAI_KHOAN_DOI",
                                 "LIEN_KET_DANG_KY_DOI"]
                sheet_ids={tab:sheet(tab).id for tab in sheetnames}
                requests=team_change_requests(
                    live_plan,sheet_ids,now(),auth()["username"],
                    account_rows=account_rows,active_links=link_rows)
                client()[0].batch_update({"requests":requests})
                read_sheet.clear()
                st.session_state["team_change_notice"]=(
                    f"Đã xử lý {len(live_plan['changes'])} nhiệm vụ cho đội "
                    f"{live_plan['team_id']}; giữ {len(live_plan['protected'])} "
                    "nhiệm vụ/chứng từ để đối soát. KHO_PHAN_BO không thay đổi.")
                st.rerun()
            except Exception as exc:
                read_sheet.clear()
                st.error("Chưa xác nhận được kết quả thu hồi ("+
                         type(exc).__name__+"). KHÔNG bấm lại ngay; "
                         "ADMIN cần đối soát các sheet VAN_CHUYEN, LAP_DAT.")



def teams_and_approval():
    require_role("ADMIN")
    st.subheader("👥 Quản lý đội / Duyệt đăng ký")
    if st.button("🔄 LÀM MỚI DỮ LIỆU TỪ GOOGLE SHEETS",
                 key="admin_team_refresh",use_container_width=True):
        read_sheet.clear()
        st.rerun()
    registrations=read_sheet("DANG_KY_THANH_VIEN")
    active_teams=[t for t in read_sheet("QUAN_LY_DOI")
                  if t.get("Mã đội") and t.get("Trạng thái")=="Hoạt động"]
    account_count=sum(bool(r.get("Mã tài khoản"))
                      for r in read_sheet("TAI_KHOAN_DOI")
                      if r.get("Trạng thái")=="Hoạt động")
    st.caption(f"Google Sheets DỰ ÁN: **{len(registrations)} hồ sơ đăng ký**"
               f" • **{len(active_teams)} đội trưởng đang hoạt động**"
               f" • **{account_count} tài khoản cấp qua app**.")
    if not registrations and active_teams:
        st.info("📋 Bảng DANG_KY_THANH_VIEN đã trống. "
                "Các tên trong mục CẤP / ĐẶT LẠI MẬT KHẨU là đội trưởng "
                "đang còn ở bảng QUAN_LY_DOI, KHÔNG phải hồ sơ chờ duyệt. "
                "Các đội cũ/tạo thủ công chưa có liên kết hồ sơ "
                "sẽ được giữ nguyên để tránh xóa nhầm. "
                "Với đội mới có liên kết, khi xóa hồ sơ app sẽ tự ngừng "
                "tài khoản và gỡ các nhiệm vụ CHƯA thực hiện lúc ADMIN "
                "mở/làm mới app. Phân bổ kho cho xã vẫn giữ nguyên.")
    render_pending_approvals(
        read_sheet=read_sheet, sheet=sheet, client=client, auth=auth,
        audit=audit, lookup=lookup, sheet_datetime_serial=sheet_datetime_serial,
        secret_accounts=st.secrets.get("accounts",{}),current_time=now)
    # Read-only impact preview: assess consequence BEFORE deleting an
    # approved registration in Sheets; no task/stock/account writes.
    links=[x for x in read_sheet("LIEN_KET_DANG_KY_DOI")
           if x.get("Trạng thái liên kết")==LINK_ACTIVE]
    with st.expander("🔎 XEM ẢNH HƯỞNG KHI XÓA ĐỘI TRƯỞNG",
                     expanded=False):
        st.caption("Chỉ mô phỏng bằng dữ liệu hiện tại, KHÔNG xóa. "
                   "ADMIN thấy nhiệm vụ nào có thể giao lại ngay, "
                   "nhiệm vụ nào phải bàn giao thực tế/đối soát trước. "
                   "Hàng phân bổ cho xã luôn giữ nguyên.")
        if not links:
            st.info("Chưa có đội nào gắn được với hồ sơ đăng ký. "
                    "Các đội nhập thủ công/cũ phải đối chiếu riêng.")
        else:
            all_teams=read_sheet("QUAN_LY_DOI")
            by_code={(t.get("Mã dự án"),t.get("Mã đội")):t for t in all_teams}
            selected=st.selectbox(
                "Chọn đội trưởng để xem tác động",
                links,
                format_func=lambda link:(
                    by_code.get((link.get("Mã dự án"),link.get("Mã đội")),{})
                    .get("Tên đội/đội trưởng","Đội trưởng")+
                    " • "+str(link.get("Mã dự án",""))),
                key="admin_detach_impact_team")
            try:
                impact=plan_detach(
                    selected,all_teams,read_sheet("TAI_KHOAN_DOI"),
                    read_sheet("VAN_CHUYEN"),read_sheet("LAP_DAT"),
                    read_sheet("BAO_CAO_TRIEN_KHAI"),
                    read_sheet("BAN_GIAO_TAP_KET"))
                c1,c2=st.columns(2)
                c1.metric("SKU-nhiệm vụ trả chờ giao",
                          len(impact["pending_tasks"]))
                c2.metric("Nhiệm vụ/chứng từ giữ đối soát",
                          len(impact["protected_tasks"]))
                st.caption(f"Địa điểm có việc trả về: "
                           f"{impact['released_sites']} • "
                           f"KTV còn giữ hàng tập kết: "
                           f"{impact['custody_sites']} • "
                           f"VC đã bàn giao ở tập kết: "
                           f"{impact['vc_hub_sites']} • "
                           f"Biên bản cũ thiếu mã VC: "
                           f"{impact['unknown_hub_sites']}")
                if impact["protected_details"]:
                    st.warning("Có hàng đã nhận, báo cáo hoặc chứng từ "
                               "cần xử lý bàn giao. Không tự giải phóng "
                               "trách nhiệm KTV đang giữ hàng.")
                display=impact["pending_details"]+impact["protected_details"]
                if display:
                    display_centered_table([{
                        "Công đoạn":r["table"],"Mã điểm":r["point_id"],
                        "SKU":r["sku"],"SL nhiệm vụ":r["quantity"],
                        "Xử lý":r["state"],
                        "Lý do":r.get("reason","")
                    } for r in display],
                        hide_index=True,use_container_width=True,height=230)
                else:
                    st.info("Đội chưa có nhiệm vụ VC/LĐ cần xử lý.")
                st.caption("Nếu V xóa hồ sơ ở Google Sheets, ADMIN "
                           "mở/làm mới app sẽ đối soát theo bảng liên kết; "
                           "báo cáo đã hoàn tất và KHO_PHAN_BO giữ nguyên.")
            except ValueError as exc:
                st.error("Liên kết hoặc đội chưa hợp lệ: "+str(exc))
    provision_team_password()
    supplemental_site_assignment_admin()
    team_reassignment_admin()
    st.caption("Chỉ quản lý đội trưởng. Thành viên của đội do đội trưởng tự quản lý.")
    st.info("Địa điểm được cấp ngay khi ADMIN duyệt đăng ký đội trưởng. "
            "Muốn tăng tiến độ: thu hồi bớt các điểm chưa làm ở đội cũ, "
            "sau đó duyệt các điểm đó cho đội mới. Không cần giao địa bàn lần hai.")
    return
    with st.expander("➕ Thêm đội trưởng thủ công (không phải gán nhiệm vụ)", expanded=False):
        saved_team_message=st.session_state.pop("team_saved_message", None)
        if saved_team_message:
            st.success(saved_team_message)
        revision=st.session_state.get("team_form_revision", 0)
        with st.form(f"team_form_{revision}"):
            tid=st.text_input("Mã đội mới", key=f"team_id_{revision}")
            name=st.text_input("Tên đội / Đội trưởng", key=f"team_name_{revision}")
            phone=st.text_input("SĐT đội trưởng", key=f"team_phone_{revision}")
            kind=st.selectbox("Loại nhiệm vụ", ["— Chọn loại nhiệm vụ —","VC","KTV","VC+KTV"],
                              key=f"team_kind_{revision}")
            proj=[r["Mã dự án"] for r in read_sheet("DANH_SACH_DU_AN") if r.get("Mã dự án")]
            project=st.selectbox("Dự án", ["— Chọn dự án —"]+proj,
                                 key=f"team_project_{revision}")
            if st.form_submit_button("THÊM ĐỘI"):
                tid, name, phone=tid.strip(), name.strip(), phone.strip()
                if not tid or not name or not phone or kind not in ("VC","KTV","VC+KTV") or project not in proj:
                    st.error("Vui lòng nhập đủ mã đội, tên, số điện thoại và chọn nhiệm vụ, dự án.")
                elif lookup(read_sheet("QUAN_LY_DOI"),"Mã đội",tid):
                    st.error("Mã đội này đã tồn tại; hãy chọn mã khác.")
                elif matching_teams(read_sheet("QUAN_LY_DOI"),name=name,phone=phone,
                                    kind=kind,project=project):
                    st.error("Đã có đội trùng người phụ trách, số điện thoại và dự án.")
                else:
                    try:
                        append("QUAN_LY_DOI",[tid,name,phone,kind,"","Hoạt động",project,""])
                        st.session_state["team_saved_message"]=(
                            f"Đã thêm đội {tid} thành công. Biểu mẫu đã được làm trống "
                            "để nhập đội tiếp theo; đội chưa được cấp tài khoản báo cáo.")
                        st.session_state["team_form_revision"]=revision+1
                        st.rerun()
                    except Exception as exc:
                        st.error("Không lưu được đội ("+type(exc).__name__+
                                 "). Dữ liệu nhập vẫn được giữ nguyên để kiểm tra.")
    st.markdown("#### 📍 GIAO ĐỊA BÀN CHO ĐỘI TRƯỞNG")
    assign_mode=st.radio("Phương thức phân công",
        ["Theo nhiều địa điểm (gõ tìm kiếm)", "Theo từng SKU (nâng cao)"],
        key="preview_assign_mode", horizontal=False)
    if assign_mode.startswith("Theo nhiều"):
        st.caption("ADMIN chỉ chọn đội trưởng và các địa điểm theo hợp đồng; "
                   "hệ thống tự lấy đủ các SKU và số lượng đã phân bổ. "
                   "Có thể gõ tên xã, đơn vị hoặc mã điểm để tìm ngay.")
        saved_notice=st.session_state.pop("assignment_site_notice",None)
        if saved_notice:
            st.success(saved_notice)
        all_projects=[r["Mã dự án"] for r in read_sheet("DANH_SACH_DU_AN")
                      if r.get("Mã dự án")]
        if not all_projects:
            st.info("Cần tạo dự án trước khi giao địa bàn.")
            return
        project=st.selectbox("1. Dự án",all_projects,key="site_assign_project")
        teams=[t for t in read_sheet("QUAN_LY_DOI") if
               t.get("Mã dự án")==project and t.get("Trạng thái")=="Hoạt động"
               and allowed_modes(t)]
        if not teams:
            st.warning("Dự án chưa có đội trưởng VC/LĐ đã duyệt.")
            return
        selected_team=st.selectbox("2. Đội trưởng nhận địa bàn",teams,
            format_func=lambda t:(t.get("Tên đội/đội trưởng","")+
                                  " • "+t.get("Loại nhiệm vụ","")),
            key=f"site_assign_team_{project}")
        mode_labels={"VC":"Vận chuyển (VC)", "LD":"Lắp đặt (LĐ)",
                     "VC+LD":"Cả vận chuyển và lắp đặt"}
        choices=allowed_modes(selected_team)
        role_mode=st.radio("3. Loại công việc được giao",
            choices,format_func=lambda m:mode_labels[m],
            key=f"site_assign_role_{project}_{selected_team['Mã đội']}",
            horizontal=len(choices)>1)
        kinds=("VAN_CHUYEN","LAP_DAT") if role_mode=="VC+LD" else (
            ("VAN_CHUYEN",) if role_mode=="VC" else ("LAP_DAT",))
        tasksets={tab:read_sheet(tab) for tab in kinds}
        places=read_sheet("DANH_SACH_DIEM")
        try:
            # All project locations remain searchable, including places without
            # stock yet. Eligibility is checked BEFORE confirmation.
            catalog=project_site_catalog(selected_team,role_mode,places,tasksets)
        except ValueError as exc:
            st.error(str(exc))
            return
        eligible={pid:(loc,counts) for pid,(loc,status,counts) in catalog.items()
                  if status=="Sẵn sàng giao"}
        st.caption(f"**{len(catalog)} địa điểm trong dự án** • "
                   f"{len(eligible)} điểm có nhiệm vụ chờ gán. "
                   "Có thể tìm toàn bộ xã/đơn vị; điểm chưa phân bổ hoặc "
                   "đã gán đội sẽ hiện trạng thái rõ ràng.")
        if not catalog:
            st.warning("Dự án chưa có danh sách địa điểm. "
                       "Admin cần nạp DANH_SACH_DIEM trước.")
            return
        def site_name(pid):
            loc,status,_=catalog[pid]
            name=loc.get("Tên xã/phường/đơn vị","") or pid
            return f"{name} • {pid} — {status}"
        # Sort ready sites first, then the rest; do not hide project locations.
        options=sorted(catalog,key=lambda pid:(
            catalog[pid][1]!="Sẵn sàng giao",
            (catalog[pid][0].get("Tên xã/phường/đơn vị","") or pid).casefold()))
        chosen=st.multiselect(
            "4. 🔎 Tìm và chọn nhiều địa điểm theo hợp đồng",
            options,format_func=site_name,
            placeholder="Gõ tên xã, đơn vị hoặc mã điểm để tìm...",
            key=f"site_assign_multi_{project}_{selected_team['Mã đội']}_{role_mode}")
        st.caption("Gõ tên hoặc mã điểm để lọc gợi ý. "
                   "Danh sách không biến mất khi chọn xong một điểm. "
                   "Chỉ điểm 'Sẵn sàng giao' mới được xác nhận phân công.")
        if not chosen:
            st.info("Chọn các địa điểm trong hợp đồng để kiểm tra tình trạng giao việc.")
            return
        ineligible=[pid for pid in chosen if pid not in eligible]
        st.write(f"**Đã chọn {len(chosen)} / {len(catalog)} địa điểm** "
                 f"({len(chosen)-len(ineligible)} sẵn sàng giao).")
        if ineligible:
            st.warning("Có địa điểm chưa đủ điều kiện, chưa thể xác nhận. "
                       "Vẫn có thể tìm trong danh sách nhưng cần phân bổ hàng "
                       "trước hoặc kiểm tra đội đã được giao.")
            for pid in ineligible[:15]:
                st.write("• "+site_name(pid))
            if len(ineligible)>15:
                st.caption(f"Còn {len(ineligible)-15} điểm chưa đủ điều kiện.")
            return
        try:
            planned=site_assignment_plan(
                selected_team,role_mode,chosen,places,tasksets)
        except ValueError as exc:
            st.error(str(exc))
            return
        st.markdown("#### 5. Kiểm tra địa bàn và nhiệm vụ trước khi giao")
        counts={tab:len(rows) for tab,rows in planned.items()}
        st.write(f"**{len(chosen)} địa điểm** · "
                 f"**{sum(counts.values())} nhiệm vụ/SKU** sẽ được gán một lần.")
        summary=[]
        for pid in chosen:
            loc, by_role=eligible[pid]
            summary.append({
                "Địa điểm":loc.get("Tên xã/phường/đơn vị",""),
                "Mã điểm":pid,
                "Nhiệm vụ VC":by_role.get("VAN_CHUYEN",0),
                "Nhiệm vụ LĐ":by_role.get("LAP_DAT",0)
            })
        display_centered_table(summary,hide_index=True,use_container_width=True)
        with st.expander("Xem SKU và số lượng tại các địa điểm đã chọn"):
            for tab,rows in planned.items():
                st.write("**"+("Vận chuyển" if tab=="VAN_CHUYEN" else "Lắp đặt")+"**")
                display_centered_table([{
                    "Điểm":r.get("Mã điểm"),"SKU":r.get("SKU"),
                    "Số lượng":r.get("SL giao" if tab=="VAN_CHUYEN" else "SL lắp")
                } for r in rows],hide_index=True,use_container_width=True)
        review_key="site_contract_checked_"+hashlib.sha256(
            repr((project,selected_team["Mã đội"],role_mode,tuple(sorted(chosen)),
                  tuple(sorted((tab,r.get("Mã nhiệm vụ VC" if tab=="VAN_CHUYEN"
                         else "Mã nhiệm vụ LD")) for tab,rr in planned.items()
                         for r in rr)))).encode("utf-8")).hexdigest()[:14]
        checked=st.checkbox("Tôi đã đối chiếu các địa điểm được giao với hợp đồng.",
                            key=review_key)
        if st.button("✅ XÁC NHẬN GIAO CÁC ĐỊA ĐIỂM ĐÃ CHỌN",
                     type="primary",use_container_width=True,
                     disabled=not checked,
                     key="site_assignment_confirm"):
            try:
                # Re-read the current spreadsheet; reject stale selections.
                read_sheet.clear()
                live_teams=read_sheet("QUAN_LY_DOI")
                live_team=lookup(live_teams,"Mã đội",selected_team["Mã đội"])
                if not live_team:
                    st.error("Đội đã bị xóa hoặc thay đổi.")
                    return
                live_places=read_sheet("DANH_SACH_DIEM")
                live_tasks={tab:read_sheet(tab) for tab in kinds}
                actual=site_assignment_plan(live_team,role_mode,
                    chosen,live_places,live_tasks)
                taskkeys={"VAN_CHUYEN":"Mã nhiệm vụ VC",
                          "LAP_DAT":"Mã nhiệm vụ LD"}
                if any({r[taskkeys[tab]] for r in actual.get(tab,[])} !=
                       {r[taskkeys[tab]] for r in planned.get(tab,[])}
                       for tab in kinds):
                    st.error("Phân bổ/phân công đã thay đổi. Chưa ghi; "
                             "vui lòng đối soát rồi xác nhận lại.")
                    return
                sheetids={tab:sheet(tab).id for tab in kinds}
                updates=compact_site_assignment_requests(
                    sheetids,actual,live_team["Mã đội"])
                client()[0].batch_update({"requests":updates})
                read_sheet.clear()
                st.session_state["assignment_site_notice"]=(
                    f"Đã giao {len(chosen)} địa điểm, {sum(len(x) for x in actual.values())} "
                    f"nhiệm vụ cho đội trưởng {live_team['Tên đội/đội trưởng']}. "
                    "Không thay đổi số lượng hoặc phân bổ kho.")
                st.rerun()
            except Exception as exc:
                read_sheet.clear()
                st.error("Chưa xác nhận được kết quả ghi ("+type(exc).__name__+
                         "). KHÔNG bấm lại ngay; cần kiểm tra VAN_CHUYEN "
                         "và LAP_DAT.")
        return
    st.caption("Chế độ nâng cao: gán một SKU hoặc toàn bộ SKU của một điểm. "
               "Chỉ sử dụng khi cần phân chia nhiệm vụ cho nhiều đội.")
    tab=st.selectbox("Loại nhiệm vụ cần gán",["VAN_CHUYEN","LAP_DAT"])
    task_key = "Mã nhiệm vụ VC" if tab=="VAN_CHUYEN" else "Mã nhiệm vụ LD"
    team_key = "Mã đội VC" if tab=="VAN_CHUYEN" else "Mã đội KTV"
    required_kinds = ("VC","VC+KTV") if tab=="VAN_CHUYEN" else ("KTV","VC+KTV")
    waiting_status = "Chưa giao" if tab=="VAN_CHUYEN" else "Chưa lắp"
    all_tasks=read_sheet(tab)
    tasks=[x for x in all_tasks if x.get(task_key) and not x.get(team_key)
           and x.get("Trạng thái")==waiting_status]
    if not tasks:
        st.info("Chưa có nhiệm vụ đang chờ phân công.")
        return
    allocation_names={r.get("Mã phân bổ"):r.get("Tên hàng hóa","")
                      for r in read_sheet("KHO_PHAN_BO")
                      if r.get("Mã phân bổ")}
    place_names={r.get("Mã điểm"):r.get("Tên xã/phường/đơn vị","")
                 for r in read_sheet("DANH_SACH_DIEM")
                 if r.get("Mã dự án")==tasks[0].get("Mã dự án")}
    selected_task=st.selectbox("Công việc chưa gán",tasks,
        format_func=lambda x:(
            f"{x['Mã dự án']} • "
            f"{place_names.get(x.get('Mã điểm')) or x.get('Mã điểm')} "
            f"({x.get('Mã điểm')}) • "
            f"{allocation_names.get(x.get('Mã phân bổ')) or x.get('SKU')} "
            f"(SKU {x.get('SKU')}) • {x.get(task_key,'')}"))
    teams=[x for x in read_sheet("QUAN_LY_DOI") if
           x.get("Trạng thái")=="Hoạt động"
           and x.get("Mã dự án")==selected_task.get("Mã dự án")
           and x.get("Loại nhiệm vụ") in required_kinds]
    if not teams:
        st.warning("Chưa có đội hoạt động đúng chuyên môn TRONG CÙNG DỰ ÁN.")
        return
    selected_team=st.selectbox("Đội trưởng nhận nhiệm vụ",teams,
        format_func=lambda x:
            f"{x.get('Tên đội/đội trưởng') or 'Chưa có tên'} • {x.get('Mã đội')}")
    group_site=st.checkbox("Gán tất cả SKU chưa gán tại điểm này cho cùng đội",
                           value=True, key=f"assign_entire_point_{tab}")
    try:
        preview=assignment_rows(
            all_tasks,selected_task,selected_team,task_key,team_key,
            required_kinds,waiting_status,group_site)
    except ValueError as exc:
        st.error(str(exc))
        return
    st.write(f"**Dự kiến gán {len(preview)} dòng nhiệm vụ** cho "
             f"{selected_team['Tên đội/đội trưởng']} tại "
             f"{selected_task['Mã điểm']}.")
    with st.expander("Xem các SKU và số lượng sẽ giao"):
        for task in preview:
            item_name=allocation_names.get(task.get("Mã phân bổ")) or task.get("SKU")
            st.write(f"{item_name} • SKU {task['SKU']}: "
                     f"{task.get('SL giao' if tab=='VAN_CHUYEN' else 'SL lắp','')} "
                     f"({task[task_key]})")
    if st.button("✅ XÁC NHẬN GÁN ĐỘI", use_container_width=True):
        try:
            # Re-read live data to reject stale UI, other projects or wrong role.
            read_sheet.clear()
            latest_tasks=read_sheet(tab)
            latest_team=lookup(read_sheet("QUAN_LY_DOI"),"Mã đội",
                               selected_team["Mã đội"])
            if not latest_team:
                st.error("Đội không còn tồn tại.")
                return
            selected_live=lookup(latest_tasks,task_key,selected_task[task_key])
            if not selected_live:
                st.error("Không tìm thấy nhiệm vụ được chọn.")
                return
            now_rows=assignment_rows(
                latest_tasks,selected_live,latest_team,task_key,team_key,
                required_kinds,waiting_status,group_site)
            if {x[task_key] for x in now_rows}!={x[task_key] for x in preview}:
                st.error("Danh sách nhiệm vụ đã thay đổi; không gán. "
                         "Vui lòng kiểm tra và xác nhận lại.")
                return
            batch=assignment_batch_requests(sheet(tab).id,now_rows,
                                            latest_team["Mã đội"])
            client()[0].batch_update({"requests":batch})
            read_sheet.clear()
            st.success(f"Đã gán {len(now_rows)} dòng cho đội "
                       f"{latest_team['Mã đội']} tại {selected_task['Mã điểm']}. "
                       "Chỉ đội trưởng được cấp tài khoản mới báo cáo được.")
            st.rerun()
        except Exception as exc:
            read_sheet.clear()
            st.error("Chưa xác nhận được kết quả ghi ("+type(exc).__name__+
                     "). Không bấm lại ngay; cần kiểm tra sheet "+tab+".")

def sync_progress(project_id, point_id):
    """Legacy ADMIN reconciliation; use the same KTV rules as atomic reports."""
    read_sheet.clear()
    summary=site_progress(project_id,point_id,
                          read_sheet("VAN_CHUYEN"),read_sheet("LAP_DAT"))
    loc=next((r for r in read_sheet("DANH_SACH_DIEM")
              if r.get("Mã dự án")==project_id and r.get("Mã điểm")==point_id),{})
    current=next((r for r in read_sheet("TIEN_DO")
                  if r.get("Mã dự án")==project_id and r.get("Mã điểm")==point_id),None)
    values=[project_id,point_id,loc.get("Tên xã/phường/đơn vị",""),
            summary["total_vc"],summary["vc_done"],
            summary["total_ld"],summary["ld_done"],summary["status"],now()]
    if current:
        sheet("TIEN_DO").update(f"A{current['_row']}:I{current['_row']}",
                               [values],value_input_option="RAW")
    else:
        append("TIEN_DO",values)
    read_sheet.clear()


def images_to_drive(files, report_id, project_id="", point_id="", point_name="", recorded_at=""):
    folder=st.secrets.get("photos_folder_id", "")
    if not folder:
        raise RuntimeError("Thiếu photos_folder_id.")
    drive=drive_client()
    # Searchable Drive filename/description; the report Sheet stores the URL.
    if project_id and point_id and recorded_at:
        safe=lambda x: "".join(c for c in str(x) if c.isalnum() or c in "-_")[:48]
        stamp="".join(c for c in recorded_at if c.isdigit() or c=="T")[:15]
        photo_prefix=f"{safe(project_id)}_{safe(point_id)}_{stamp}_{report_id}"
        description=(f"TDV DỰ ÁN | Dự án: {project_id} | "
                     f"Điểm: {point_name} ({point_id}) | "
                     f"Thời điểm gửi: {recorded_at} | Mã báo cáo: {report_id}")
    else:
        photo_prefix=report_id
        description=f"TDV DỰ ÁN | Mã chứng từ: {report_id}"
    urls=[]
    for i, f in enumerate(files[:10]):
        img=Image.open(f)
        img=img.convert("RGB")
        img.thumbnail((1280,1280))
        buf=io.BytesIO()
        img.save(buf,format="JPEG",quality=77,optimize=True)
        buf.seek(0)
        uploaded=drive.files().create(
            body={"name":f"{photo_prefix}_{i+1}.jpg",
                  "parents":[folder],"description":description},
            media_body=MediaIoBaseUpload(buf,mimetype="image/jpeg",resumable=False),
            supportsAllDrives=True,
            fields="id,webViewLink").execute()
        urls.append(uploaded.get("webViewLink") or
                    f"https://drive.google.com/file/d/{uploaded['id']}/view")
    return urls

def hub_admin():
    """ADMIN records ONE signed physical transfer to a KTV receiving hub.

    Site quota remains in KHO_PHAN_BO; this ledger does not write 'Đã giao'
    to each commune or silently allocate goods a second time.
    """
    require_role("ADMIN")
    st.subheader("📦 Bàn giao hàng TẬP KẾT cho KTV")
    st.caption("Ví dụ: đội VC giao gộp hàng của Đồng Văn – Yên Minh – "
               "Quản Bạ tại một điểm; KTV ký biên bản giấy nhận đủ, "
               "sau đó tổ chức đi lắp từng xã. KHÔNG tính là giao tận xã.")
    projects=[p["Mã dự án"] for p in read_sheet("DANH_SACH_DU_AN")
              if p.get("Mã dự án")]
    if not projects:
        st.info("Chưa có dự án.")
        return
    if st.button("🔄 Làm mới dữ liệu tập kết",key="refresh_hub"):
        read_sheet.clear()
        st.rerun()
    with st.expander("1. Chọn dự án, đội KTV và các xã thuộc hợp đồng",
                     expanded=True):
        project=st.selectbox("Dự án",projects,index=projects.index("DA880")
                             if "DA880" in projects else 0,key="hub_project")
        teams=[t for t in read_sheet("QUAN_LY_DOI")
               if t.get("Mã dự án")==project and
               t.get("Trạng thái")=="Hoạt động" and
               t.get("Loại nhiệm vụ") in ("KTV","VC+KTV")]
        if not teams:
            st.info("Chưa có đội KTV hoạt động trong dự án. "
                    "Hãy duyệt/gán đội trước khi lập biên bản tập kết.")
            return
        team=st.selectbox("Đội KTV nhận hàng",teams,
                          format_func=lambda x:x.get("Tên đội/đội trưởng",""),
                          key="hub_ktv_team")
        vc_teams=[t for t in read_sheet("QUAN_LY_DOI")
                  if t.get("Mã dự án")==project and
                  t.get("Trạng thái")=="Hoạt động" and
                  t.get("Loại nhiệm vụ") in ("VC","VC+KTV")]
        if not vc_teams:
            st.warning("Chưa có đội VC hoạt động trong dự án. "
                       "Cần gán đội VC trước khi bàn giao tập kết.")
            return
        vc_team=st.selectbox(
            "Đội VC đã giao hàng (theo phân công)",
            vc_teams,
            format_func=lambda x:x.get("Tên đội/đội trưởng",""),
            key="hub_vc_team")
        ld_rows=read_sheet("LAP_DAT")
        points=[p for p in read_sheet("DANH_SACH_DIEM") if
                p.get("Mã dự án")==project and
                any(r.get("Mã đội KTV")==team["Mã đội"] and
                    r.get("Mã dự án")==project and
                    r.get("Mã điểm")==p.get("Mã điểm") for r in ld_rows)]
        if not points:
            waiting={r.get("Mã điểm") for r in ld_rows
                     if r.get("Mã dự án")==project and r.get("Mã điểm")
                     and r.get("Trạng thái")=="Chưa lắp"
                     and not r.get("Mã đội KTV")}
            st.warning(
                f"Đội KTV **{team.get('Tên đội/đội trưởng','')}** "
                f"chưa được giao xã/đơn vị nào trong dự án **{project}**. "
                f"Dự án hiện có {len(waiting)} địa điểm còn nhiệm vụ "
                "LĐ chưa gán đội. Đây không phải lỗi phân bổ hàng.")
            st.info("Thứ tự test: **Quản lý đội / Duyệt đăng ký → "
                    "GIAO ĐỊA BÀN CHO ĐỘI TRƯỞNG → chọn dự án → "
                    "chọn đội KTV → chọn Lắp đặt (LĐ) → "
                    "chọn xã → xem bảng SKU → xác nhận giao nhiệm vụ.** "
                    "Sau đó thực hiện tương tự với đội VC trên đúng các xã, "
                    "rồi quay lại Bàn giao tập kết.")
            if st.button("📍 ĐẾN MỤC GIAO ĐỊA BÀN",
                         use_container_width=True,
                         key="hub_go_assignments"):
                st.session_state["_next_menu"]="👥 Đăng ký & quản lý đội"
                st.session_state["_next_assignment_project"]=project
                st.rerun()
            return
        lookup_names={p["Mã điểm"]:p.get("Tên xã/phường/đơn vị","")
                      for p in points}
        options=sorted(lookup_names,key=lambda k:lookup_names[k].casefold())
        selection_key=f"hub_selected_sites_{project}_{team['Mã đội']}"
        if st.button("☑️ CHỌN TẤT CẢ XÃ ĐƯỢC GIAO",
                     key=f"hub_select_all_{project}_{team['Mã đội']}",
                     use_container_width=True):
            st.session_state[selection_key]=list(options)
        selected=st.multiselect(
            "🔎 Tìm và chọn các xã trong chuyến bàn giao",
            options,format_func=lambda k:lookup_names[k],
            placeholder="Gõ tên xã; chọn một hoặc nhiều điểm...",
            key=selection_key
        )
        if not selected:
            st.info("Chọn các xã theo hợp đồng của đội trưởng KTV.")
            return
        try:
            lines,summary=hub_plan(team,selected,ld_rows,points,
                                   read_sheet("KHO_PHAN_BO"))
        except ValueError as exc:
            st.error(str(exc))
            return
        try:
            verify_vc_assignment(
                lines,vc_team["Mã đội"],read_sheet("VAN_CHUYEN"))
        except ValueError as exc:
            st.error("Chưa thể bàn giao tập kết: "+str(exc))
            return
        prior=read_sheet("BAN_GIAO_TAP_KET")
        accepted={(r.get("Mã dự án"),r.get("Mã điểm"),r.get("SKU"))
                  for r in prior if r.get("Trạng thái")=="KTV đã nhận tại tập kết"}
        repeats=[l for l in lines if
                 (project,l["point_id"],l["sku"]) in accepted]
        st.write(f"**{len(selected)} địa điểm · {len(lines)} dòng SKU** "
                 "thuộc biên bản tập kết.")
        if repeats:
            st.error(f"Đã có {len(repeats)} dòng SKU thuộc biên bản KTV "
                     "đã nhận tại điểm tập kết. Chưa thể ghi đúp.")
            return
    with st.expander("2. Đối chiếu đủ số hàng trong chuyến giao",
                     expanded=False):
        st.caption("Số thực nhận chỉ xác nhận khi KTV đã kiểm đếm và hai "
                   "bên đã ký biên bản giấy. Trường hợp thiếu hàng hãy "
                   "tách chuyến/xã và báo ADMIN đối soát.")
        qty_ok=True
        for (sku,unit),amount in sorted(summary.items()):
            actual=st.number_input(f"{sku} ({unit}) • phải nhận {qty_label(amount)}",
                min_value=0., value=float(amount),step=1.,
                key=f"hub_qty_{project}_{team['Mã đội']}_{sku}")
            if abs(actual-float(amount))>1e-9:
                qty_ok=False
        if not qty_ok:
            st.warning("SL thực nhận khác số dự kiến. Chưa xác nhận được "
                       "toàn bộ điểm; hãy kiểm tra biên bản/chọn lại các xã.")
        st.caption("Danh sách chi tiết nằm trong bảng cuộn; không cần mở "
                   "thêm nhiều lớp biểu mẫu trên điện thoại.")
        display_centered_table([{"Xã/đơn vị":r["name"],"SKU":r["sku"],
                       "Số lượng":qty_label(r["quantity"]),"ĐVT":r["unit"]}
                      for r in lines],hide_index=True,use_container_width=True,
                     height=180)
    with st.expander("3. Điểm tập kết và BIÊN BẢN GIẤY có ký nhận",
                     expanded=False):
        area=st.text_input("Khu vực hợp đồng (địa danh cũ)",
                           placeholder="Ví dụ: Đồng Văn, Yên Minh, Quản Bạ",
                           key="hub_area")
        hub=st.text_input("Tên điểm tập kết",placeholder="Kho/địa điểm nhận hàng",
                          key="hub_name")
        address=st.text_input("Địa chỉ điểm tập kết",
                              key="hub_address")
        giver=st.text_input(
            "Họ tên người/đội VC bàn giao trên biên bản",
            value=vc_team.get("Tên đội/đội trưởng", ""),
            key=f"hub_giver_{project}_{vc_team['Mã đội']}"
        )
        paper=st.text_input("Số biên bản bàn giao giấy",
                            key="hub_paper")
        proof=st.file_uploader(
            "📷 Ảnh biên bản giấy có chữ ký của bên giao và KTV nhận",
            type=["jpg","jpeg","png"],key="hub_scan")
        acknowledged=st.checkbox(
            "Tôi đã đối chiếu hàng THỰC NHẬN và thấy biên bản giấy có "
            "đủ chữ ký hai bên; ghi nhận đây là giao tại điểm tập kết, "
            "KHÔNG phải giao tận xã.",
            key="hub_confirm_paper")
        ready=bool(qty_ok and area.strip() and hub.strip()
                   and address.strip() and giver.strip() and paper.strip()
                   and proof and acknowledged)
        if st.button("✅ GHI NHẬN BÀN GIAO TẬP KẾT",
                     type="primary",use_container_width=True,
                     disabled=not ready,key="hub_commit"):
            try:
                read_sheet.clear()
                live_team=lookup(read_sheet("QUAN_LY_DOI"),
                                 "Mã đội",team["Mã đội"])
                if not live_team or live_team.get("Mã dự án")!=project or (
                    live_team.get("Trạng thái")!="Hoạt động"
                    or live_team.get("Loại nhiệm vụ") not in ("KTV","VC+KTV")
                ):
                    fail("Đội KTV đã thay đổi; chưa ghi dữ liệu.")
                actual,actual_sum=hub_plan(
                    live_team,selected,read_sheet("LAP_DAT"),
                    read_sheet("DANH_SACH_DIEM"),
                    read_sheet("KHO_PHAN_BO"))
                live_vc=lookup(read_sheet("QUAN_LY_DOI"),
                               "Mã đội",vc_team["Mã đội"])
                if (not live_vc or live_vc.get("Mã dự án")!=project
                        or live_vc.get("Trạng thái")!="Hoạt động"
                        or live_vc.get("Loại nhiệm vụ") not in ("VC","VC+KTV")):
                    fail("Đội VC đã thay đổi, chưa được ghi bàn giao.")
                verify_vc_assignment(
                    actual,vc_team["Mã đội"],read_sheet("VAN_CHUYEN"))
                current_set={(r["point_id"],r["sku"],str(r["quantity"]))
                             for r in actual}
                original={(r["point_id"],r["sku"],str(r["quantity"]))
                          for r in lines}
                if current_set!=original:
                    fail("Phân công/định mức đã thay đổi; đối chiếu rồi xác nhận lại.")
                # Prevent two independent delivery proofs for the same site.
                # A signed KTV hub handover is not a retroactive direct VC delivery.
                keys={(r["project"],r["point_id"],r["sku"]) for r in actual}
                direct=[r for r in read_sheet("VAN_CHUYEN")
                        if (r.get("Mã dự án"),r.get("Mã điểm"),r.get("SKU")) in keys
                        and r.get("Trạng thái")=="Đã giao hàng"]
                completed=[r for r in read_sheet("LAP_DAT")
                           if (r.get("Mã dự án"),r.get("Mã điểm"),r.get("SKU")) in keys
                           and r.get("Trạng thái")=="Đã lắp đặt"]
                if direct or completed:
                    fail("Một số SKU đã có xác nhận giao tận xã hoặc đã lắp xong. "
                         "Chưa được ghi trùng biên bản tập kết.")
                existed={(r.get("Mã dự án"),r.get("Mã điểm"),r.get("SKU"))
                         for r in read_sheet_live("BAN_GIAO_TAP_KET")
                         if r.get("Trạng thái")=="KTV đã nhận tại tập kết"}
                if any((project,r["point_id"],r["sku"]) in existed
                       for r in actual):
                    fail("Các SKU này đã có biên bản tập kết; không ghi đúp.")
                batch="BG-"+uuid.uuid4().hex[:12].upper()
                urls=images_to_drive([proof],batch)
                if not urls:
                    fail("Không lưu được ảnh biên bản; chưa ghi bàn giao.")
                rows=signed_receipt_lines(batch,live_team,actual,
                    area.strip(),hub.strip(),address.strip(),
                    giver.strip(),paper.strip(),urls[0],now(),
                    vc_team_id=vc_team["Mã đội"])
                sheet("BAN_GIAO_TAP_KET").append_rows(rows,
                                                      value_input_option="RAW")
                read_sheet.clear()
                audit(auth()["username"],"BÀN GIAO TẬP KẾT",
                      "BAN_GIAO_TAP_KET",batch,"OK",
                      f"{len(selected)} điểm, {len(rows)} dòng; số giấy {paper}")
                st.success("Đã ghi biên bản tập kết. KTV nhận hàng theo "
                           "từng SKU của các xã đã chọn; chưa đánh dấu "
                           "VC giao tận xã hoặc lắp đặt xong.")
                st.rerun()
            except Exception as exc:
                read_sheet.clear()
                st.error("Chưa xác nhận ghi đủ biên bản ("+
                         type(exc).__name__+"). KHÔNG bấm lại; "
                         "ADMIN cần kiểm tra BAN_GIAO_TAP_KET.")
    st.caption("Biên bản gắn MÃ ĐỘI VC giao và MÃ ĐỘI KTV nhận. "
               "Hai bên có trách nhiệm đối soát theo từng SKU. "
               "Bảng biên bản chỉ dùng trên DỰ ÁN. "
               "Không thay phiếu xuất kho hoặc biên bản nghiệm thu lắp đặt.")


def report():
    a=require_role("DOI_TRUONG","ADMIN")
    st.subheader("📋 Báo cáo nhiệm vụ")
    tid=a["team_id"]
    if a["role"]=="ADMIN":
        allteams=[x for x in read_sheet("QUAN_LY_DOI") if x.get("Mã đội")]
        if not allteams:
            st.info("Chưa có đội.")
            return
        selected_report_team=st.selectbox(
            "Kiểm thử theo đội (chỉ Admin)",allteams,
            format_func=lambda x:
                f"{x.get('Tên đội/đội trưởng') or 'Chưa có tên'} • {x.get('Mã đội')}")
        tid=selected_report_team["Mã đội"]
    if not tid:
        fail("Chưa liên kết mã đội cho tài khoản.")
    report_team=lookup(read_sheet("QUAN_LY_DOI"),"Mã đội",tid) or {}
    if report_team:
        st.caption("Đội đang báo cáo: "
                   f"{report_team.get('Tên đội/đội trưởng') or 'Chưa có tên'}"
                   f" • Mã đội: {tid}")
    if a["role"] == "ADMIN":
        with st.expander("🔐 Kết nối thư mục ảnh DỰ ÁN (chỉ thư mục đã chọn)",
                         expanded=False):
            configure_photo_folder()
    if st.button("🔄 CẬP NHẬT HÀNG VC ĐÃ GIAO",use_container_width=True,
                 key="team_report_refresh"):
        st.session_state.pop("_report_attempt_pending",None)
        read_sheet.clear()
        st.rerun()
    all_vc=read_sheet("VAN_CHUYEN")
    all_ld=read_sheet("LAP_DAT")
    hub_rows=read_sheet("BAN_GIAO_TAP_KET")
    # Site-level rights are visible before any SKU allocation.
    # This table is read-only; submitting a report still requires a task.
    team=next((r for r in read_sheet("QUAN_LY_DOI")
               if r.get("Mã đội")==tid),None)
    if team is not None:
        try:
            approved_sites=approved_team_site_overview(
                team,read_sheet("DANG_KY_THANH_VIEN"),
                read_sheet("DANH_SACH_DIEM"),
                {"VAN_CHUYEN":all_vc,"LAP_DAT":all_ld})
        except ValueError as exc:
            st.warning("Quyền địa điểm cần ADMIN đối soát: "+str(exc))
            approved_sites=[]
        if approved_sites:
            with st.expander("📍 ĐỊA ĐIỂM ADMIN ĐÃ DUYỆT",expanded=True):
                display_centered_table([{
                    "Địa điểm":r["site_name"],
                    "Công việc":r["role"],
                    "Số nhiệm vụ SKU":r["sku_tasks"],
                    "Đã liên kết":r["linked_tasks"],
                    "Trạng thái":r["status"]
                } for r in approved_sites],
                hide_index=True,use_container_width=True)
                st.caption("Được duyệt địa điểm chưa có nghĩa hàng đã phân bổ "
                           "hoặc VC/LĐ đã hoàn thành.")
    accepted_hub={
        (r.get("Mã dự án"),r.get("Mã điểm"),r.get("SKU"),
         r.get("Mã đội VC giao"))
        for r in hub_rows
        if r.get("Trạng thái")=="KTV đã nhận tại tập kết"
        and r.get("Số biên bản giấy") and r.get("Link ảnh biên bản ký")
        and r.get("Mã đội VC giao")
    }
    vc=[x for x in all_vc if x.get("Mã đội VC")==tid
        and x.get("Trạng thái")!="Đã giao hàng"
        and (x.get("Mã dự án"),x.get("Mã điểm"),x.get("SKU"),tid)
        not in accepted_hub]
    ld=[x for x in all_ld if x.get("Mã đội KTV")==tid
        and x.get("Trạng thái")!="Đã lắp đặt"]
    handover=installer_handover(tid,all_ld,all_vc,
                                read_sheet("DANH_SACH_DIEM"),hub_rows)
    status_by_site={(r["project"],r["point_id"]):r for r in handover}
    if handover:
        st.caption("Đội KTV xem nơi đã nhận hàng tập kết và các xã đã "
                   "giao trực tiếp, không nhầm hai hình thức bàn giao.")
    if handover:
        with st.expander("📦 1. KIỂM TRA HÀNG VC / TẬP KẾT",
                         expanded=True):
            st.markdown("#### Hàng VC đã giao tận xã / KTV nhận tập kết")
            st.caption("Chỉ các địa điểm được giao cho đội LĐ này. "
                       "Số liệu VC được đối chiếu theo từng SKU và SL THỰC GIAO; "
                       "phân bổ hàng chưa phải đã giao hàng.")
            c1,c2=st.columns(2)
            c1.metric("Đủ hàng để lắp",
                      sum(r["ready"] for r in handover))
            c2.metric("Chưa sẵn sàng",
                      sum(r["status"] in ("Chưa giao","Đã giao một phần",
                                          "Đã tập kết một phần","Cần đối soát")
                          for r in handover))
            labels=["Tất cả","Sẵn sàng lắp","Sẵn sàng lắp (tập kết)",
                    "Đang lắp một phần","Đã tập kết một phần","Đã giao một phần",
                    "Chưa giao","Cần đối soát","Đã lắp xong"]
            kind=st.selectbox("Lọc trạng thái hàng VC",labels,
                              key="team_handover_filter")
            shown=[r for r in handover if kind=="Tất cả" or r["status"]==kind]
            if shown:
                display_centered_table([{
                    "Địa điểm":r["name"],"Dự án":r["project"],
                    "Tình trạng hàng":r["status"],
                    "SKU phải lắp":r["sku_count"]
                } for r in shown],hide_index=True,use_container_width=True)
                detail=st.selectbox(
                    "Xem SL hàng VC đã giao tại một điểm",
                    list(range(len(shown))),
                    format_func=lambda i:shown[i]["name"]+" • "+shown[i]["status"],
                    key="team_handover_details")
                info=shown[detail]
                st.write("**"+info["name"]+"** • "+info["project"])
                signed_hubs=[r for r in hub_rows if
                    r.get("Mã dự án")==info["project"] and
                    r.get("Mã điểm")==info["point_id"] and
                    r.get("Mã đội KTV nhận")==tid and
                    r.get("Trạng thái")=="KTV đã nhận tại tập kết"
                    and r.get("Số biên bản giấy")
                    and r.get("Link ảnh biên bản ký")]
                if signed_hubs:
                    refs=sorted(set((r.get("Điểm tập kết",""),
                                     r.get("Địa chỉ tập kết",""),
                                     r.get("Khu vực theo hợp đồng",""),
                                     r.get("Số biên bản giấy",""))
                                    for r in signed_hubs))
                    for hname,haddress,area,paper in refs:
                        st.info("📦 KTV đã nhận tại **"+hname+"**"+
                                (" — "+haddress if haddress else "")+
                                (" • "+area if area else "")+
                                ". Biên bản giấy: "+paper)
                    image=signed_hubs[0].get("Link ảnh biên bản ký","")
                    if image.startswith("https://drive.google.com/"):
                        st.link_button("📄 Xem ảnh biên bản ký nhận",
                                       image,use_container_width=True)
                # Fallback Maps search is only a candidate, not a verified pin.
                dest=next((p for p in read_sheet("DANH_SACH_DIEM")
                           if p.get("Mã dự án")==info["project"] and
                           p.get("Mã điểm")==info["point_id"]),{})
                if dest.get("Vĩ độ") and dest.get("Kinh độ"):
                    query=dest["Vĩ độ"]+","+dest["Kinh độ"]
                else:
                    query=", ".join(filter(None,[
                        info["name"],dest.get("Địa chỉ chi tiết",""),
                        dest.get("Tỉnh/Thành","")]))
                    st.caption("Chưa có GPS đã xác minh: cần kiểm tra đúng đích "
                               "trên Google Maps trước khi đi.")
                st.link_button("🧭 Xem địa điểm trên Google Maps",
                    "https://www.google.com/maps/search/?api=1&query="+quote(query),
                    use_container_width=True)
                display_centered_table([{
                    "SKU":v["sku"],"Cần lắp":v["need"],
                    "VC giao xã":v["delivered"],
                    "KTV nhận tập kết":v["hub_received"],
                    "LĐ đã lắp":v["installed"],
                    "Đủ hàng":"Có" if v["delivery_ready"] else "Chưa"
                } for v in info["details"]],hide_index=True,use_container_width=True)
                for row in info["details"]:
                    if row["issues"]:
                        st.warning(row["sku"]+": "+"; ".join(row["issues"]))
            else:
                st.info("Không có địa điểm trong trạng thái vừa lọc.")
            st.divider()
    if vc and not ld:
        report_title="📝 BÁO CÁO GIAO HÀNG"
    elif ld and not vc:
        report_title="📝 BÁO CÁO LẮP ĐẶT"
    else:
        report_title="📝 BÁO CÁO GIAO HÀNG / LẮP ĐẶT"
    with st.expander(report_title, expanded=not bool(handover)):
        available=sorted(
            set((r["Mã dự án"],r["Mã điểm"]) for r in vc+ld),
            key=lambda k:(0 if status_by_site.get(k,{}).get("ready") else 1,k))
        if not available:
            st.info("Đội chưa được giao nhiệm vụ hoặc đã báo cáo hết.")
            return
        site_rows=read_sheet("DANH_SACH_DIEM")
        locs={p.get("Mã điểm"):p.get("Tên xã/phường/đơn vị")
              for p in site_rows}
        site_by_key={(p.get("Mã dự án"),p.get("Mã điểm")):p
                     for p in site_rows}
        def site_option(key):
            name=locs.get(key[1],key[1])
            summary=status_by_site.get(key)
            suffix=" • "+summary["status"] if summary else ""
            return f"{key[0]} • {name} • {key[1]}{suffix}"
        search_options=[]
        key_by_id={}
        for point_key in available:
            option_id=json.dumps(list(point_key),ensure_ascii=False,
                                 separators=(",",":"))
            search_options.append({
                "id":option_id,
                "label":site_option(point_key),
            })
            key_by_id[option_id]=point_key
        selected_key=f"report_site_selected_{tid}"
        if st.session_state.get(selected_key) not in available:
            st.session_state[selected_key]=available[0]
        search_key=f"report_site_query_{tid}"
        event_key=f"report_site_last_selection_event_{tid}"
        selected_point=st.session_state[selected_key]
        selected_id=json.dumps(list(selected_point),ensure_ascii=False,
                               separators=(",",":"))
        search_result=search_assigned_site(
            search_options,selected_id=selected_id,
            selected_label=site_option(selected_point),
            query=str(st.session_state.get(search_key) or ""),
            key=f"report_site_dropdown_{tid}")
        selection_payload=search_result
        if isinstance(selection_payload,str):
            try:
                selection_payload=json.loads(selection_payload)
            except (TypeError,ValueError):
                pass
        if isinstance(selection_payload,dict):
            result_id=str(selection_payload.get("selectedId") or "")
            event_id=str(selection_payload.get("eventId") or "")
            last_event=str(st.session_state.get(event_key) or "")
            if result_id in key_by_id and (not event_id or event_id!=last_event):
                st.session_state[selected_key]=key_by_id[result_id]
                st.session_state[search_key]=""
                st.session_state[event_key]=event_id or ("legacy:"+result_id)
                # Synchronize the chosen ID in the component props before
                # rendering the report controls below it.
                if event_id:
                    st.rerun()
        elif isinstance(selection_payload,str) and selection_payload in key_by_id:
            # Older component runtimes can return a selected JSON ID as a string.
            st.session_state[selected_key]=key_by_id[selection_payload]
            st.session_state[search_key]=""
        chosen=st.session_state[selected_key]
        site=next((p for p in read_sheet("DANH_SACH_DIEM")
                   if p.get("Mã dự án")==chosen[0]
                   and p.get("Mã điểm")==chosen[1]),{})
        if site:
            verified_pin=bool(site.get("Vĩ độ") and site.get("Kinh độ"))
            destination=(
                str(site["Vĩ độ"])+","+str(site["Kinh độ"])
                if verified_pin else ", ".join(filter(None,[
                    site.get("Tên xã/phường/đơn vị") or locs.get(chosen[1]),
                    site.get("Địa chỉ chi tiết",""),
                    site.get("Tỉnh/Thành","")]))
            )
            maps_url=("https://www.google.com/maps/dir/?api=1&destination="
                      +quote(destination))
            st.link_button("🧭 TÌM ĐƯỜNG ĐẾN ĐIỂM NÀY",
                           maps_url,use_container_width=True)
            if not verified_pin:
                st.caption("Điểm chưa có GPS xác minh; kiểm tra đúng đích "
                           "trên Google Maps trước khi bắt đầu.")
        vrows=[r for r in vc if (r["Mã dự án"],r["Mã điểm"])==chosen]
        lrows=[r for r in ld if (r["Mã dự án"],r["Mã điểm"])==chosen]
        modes=(["Đã giao hàng"] if vrows else [])+(["Đã lắp đặt"] if lrows else [])
        if not modes:
            st.info("Điểm này hiện không còn nhiệm vụ VC/KTV chờ báo cáo. "
                    "Chọn điểm khác hoặc nhờ ADMIN kiểm tra phân công.")
            return
        if vrows and lrows:
            st.caption("Đội VC+LĐ báo GIAO HÀNG trước, sau khi đã xác nhận "
                       "giao đủ mới được báo LẮP ĐẶT XONG. Hai bước lưu riêng.")
        mode=st.radio("Loại báo cáo",modes,horizontal=True)
        if mode=="Đã lắp đặt" and not site_may_complete_installation(
                chosen[0],chosen[1],lrows,all_vc,hub_rows):
            st.warning("⛔ Điểm này chưa có xác nhận VC giao tận xã HOẶC "
                       "biên bản KTV nhận tại điểm tập kết đủ tất cả SKU. "
                       "Chưa được báo LẮP XONG; cần đối soát trước.")
            return
        targets=({"VAN_CHUYEN":vrows} if mode=="Đã giao hàng" else
                 {"LAP_DAT":lrows} if mode=="Đã lắp đặt" else
                 {"VAN_CHUYEN":vrows,"LAP_DAT":lrows})
        st.write("**HÀNG HÓA ĐƯỢC PHÂN BỔ TẠI ĐIỂM**")
        st.caption("Tên hàng, đơn vị tính và số lượng lấy trực tiếp từ "
                   "KHO_PHAN_BO theo đúng dự án và địa điểm được chọn. "
                   "Chỉ xác nhận khi đã hoàn thành công việc.")
        site_allocations=[row for row in read_sheet("KHO_PHAN_BO")
                          if row.get("Mã dự án")==chosen[0]
                          and row.get("Mã điểm")==chosen[1]]
        increments={}
        sku_display=[]
        for tab,rows in targets.items():
            id_key="Mã nhiệm vụ VC" if tab=="VAN_CHUYEN" else "Mã nhiệm vụ LD"
            for r in rows:
                try:
                    item=allocated_report_item(r,site_allocations,mode)
                    remaining=item["remaining"]
                except ValueError as exc:
                    st.error("Phân bổ hàng tại điểm cần đối soát: "+str(exc))
                    return
                if remaining<=0:
                    st.error("Số lượng còn lại không hợp lệ. Cần ADMIN đối soát.")
                    return
                sku_display.append({
                    "SKU":r.get("SKU",""),
                    "Thiết bị":item["name"],
                    "SL cần " + ("giao" if tab=="VAN_CHUYEN" else "lắp"):
                        f"{remaining:g}",
                    "ĐVT":item["unit"]})
                increments[r[id_key]]=remaining
        with st.container(border=True):
            st.markdown("**📍 " + (locs.get(chosen[1]) or chosen[1]) + "**")
            st.caption(f"Dự án: {chosen[0]} • Mã điểm: {chosen[1]} • "
                       f"{len(sku_display)} SKU • "
                       f"{'Báo giao hàng' if mode == 'Đã giao hàng' else 'Báo lắp đặt'}")
            display_centered_table(
                sku_display, hide_index=True, use_container_width=True,
                center_columns=("SKU", "SL cần giao", "SL cần lắp", "ĐVT"),
                compact=True)
            if any(numeric(r.get("SL thực tế"))>0
                   for rows in targets.values() for r in rows):
                st.caption("SL ở bảng là phần còn phải thực hiện; "
                           "đã trừ số lượng báo cáo trước đó.")
        st.caption("🧪 MÃ BẢN KIỂM TRA ẢNH DRIVE: DRIVE-PICKER-APPID-20261027-B")
        photo_key=f"{chosen[0]}_{chosen[1]}_{mode}"
        uploaded=st.file_uploader(
            "📸 CHỤP HOẶC CHỌN ẢNH HIỆN TRƯỜNG — chạm để mở Camera/Thư viện "
            "(tối đa 10 ảnh)",
            type=["jpg","jpeg","png"],accept_multiple_files=True,
            key=f"report_files_{photo_key}")
        photos=list(uploaded or [])
        st.caption(f"Ảnh minh chứng đã chọn: {len(photos)}/10")
        require_gps = mode!="Đã giao hàng"
        gps=None
        if require_gps:
            st.markdown("**📍 Check-in GPS bắt buộc**")
            gps=streamlit_geolocation()
        else:
            st.caption("Vận chuyển không bắt buộc GPS check-in.")
        lat=gps.get("latitude") if isinstance(gps,dict) else None
        lon=gps.get("longitude") if isinstance(gps,dict) else None
        # Never offer an actionable submit if the PREVIEW photo destination
        # is missing; this used to surface only after clicking GỬI BÁO CÁO.
        if not str(st.secrets.get("photos_folder_id", "")).strip():
            st.error("⛔ Chưa cấu hình thư mục lưu ảnh DỰ ÁN "
                     "(photos_folder_id). Chưa thể gửi báo cáo. "
                     "ADMIN cần bổ sung cấu hình ảnh trên Streamlit Cloud; "
                     "không bấm gửi lại khi chưa hoàn tất.")
            return
        drive_oauth=st.secrets.get("google_drive_oauth", {})
        oauth_ready=all(str(drive_oauth.get(k, "")).strip()
                        for k in ("client_id", "client_secret", "refresh_token",
                                  "picker_api_key"))
        if not oauth_ready:
            st.error("⛔ Chưa cấu hình OAuth giới hạn thư mục ảnh. "
                     "ADMIN cần hoàn tất Kết nối thư mục ảnh DỰ ÁN trước; "
                     "chưa bật gửi báo cáo để tránh tạo báo cáo lỗi.")
            return
        try:
            verify_configured_photo_folder()
        except Exception as exc:
            st.error("⛔ Chưa xác minh được quyền thư mục ảnh: " +
                     drive_error_summary(exc) +
                     ". ADMIN cần chọn/xác minh thư mục trong mục "
                     "“Kết nối thư mục ảnh DỰ ÁN” trước khi gửi.")
            return
        eligible=(bool(photos) and len(photos)<=10
                  and (not require_gps or (lat is not None and lon is not None)))
        if not eligible:
            st.info("Cần ảnh và" + (" GPS check-in." if require_gps else " đúng số ảnh quy định."))
            return
        if st.button("✅ GỬI BÁO CÁO",type="primary",use_container_width=True):
            if not any(v>0 for v in increments.values()):
                st.error("Phải nhập SL thực tế vừa thực hiện lớn hơn 0.")
                return
            if st.session_state.get("_report_attempt_pending"):
                st.error("Lần gửi trước chưa được đối soát. Bấm CẬP NHẬT HÀNG VC hoặc nhờ ADMIN kiểm tra trước khi gửi lại.")
                return
            # Server-side recheck: never trust a selected team/name from browser.
            live=[]
            for tab, rows in targets.items():
                all_rows=read_sheet_live(tab)
                id_col="Mã nhiệm vụ VC" if tab=="VAN_CHUYEN" else "Mã nhiệm vụ LD"
                team_col="Mã đội VC" if tab=="VAN_CHUYEN" else "Mã đội KTV"
                closed="Đã giao hàng" if tab=="VAN_CHUYEN" else "Đã lắp đặt"
                for row in rows:
                    current=lookup(all_rows,id_col,row[id_col])
                    if (not current or current.get(team_col)!=tid or
                            current.get("Mã dự án")!=chosen[0] or
                            current.get("Mã điểm")!=chosen[1] or
                            current.get("Trạng thái")==closed or
                            current.get("Trạng thái")!=row.get("Trạng thái") or
                            numeric(current.get("SL thực tế"))!=numeric(row.get("SL thực tế"))):
                        fail("Quyền, SL thực tế hoặc trạng thái đã thay đổi. Vui lòng tải lại.")
                    live.append((tab,current))
            # Re-read live VC status before installation; no stale/cached approval.
            if mode=="Đã lắp đặt":
                read_sheet.clear()
                fresh_ld=read_sheet("LAP_DAT")
                fresh=[lookup(fresh_ld,"Mã nhiệm vụ LD",row["Mã nhiệm vụ LD"])
                       for row in lrows]
                if any(not row or row.get("Mã đội KTV")!=tid or
                       row.get("Trạng thái")=="Đã lắp đặt" for row in fresh):
                    fail("Nhiệm vụ lắp đã đổi; không báo cáo. Hãy tải lại.")
                if not site_may_complete_installation(
                        chosen[0],chosen[1],fresh,read_sheet("VAN_CHUYEN"),
                        read_sheet("BAN_GIAO_TAP_KET")):
                    fail("Hàng VC chưa giao đủ hoặc chưa đối soát được; "
                         "không thể xác nhận lắp đặt.")
            if mode=="Đã giao hàng":
                read_sheet.clear()
                current_hub=read_sheet_live("BAN_GIAO_TAP_KET")
                confirmed={(r.get("Mã dự án"),r.get("Mã điểm"),r.get("SKU"),
                            r.get("Mã đội VC giao")) for r in current_hub
                           if r.get("Trạng thái")=="KTV đã nhận tại tập kết"
                           and r.get("Số biên bản giấy")
                           and r.get("Link ảnh biên bản ký")
                           and r.get("Mã đội VC giao")}
                if any((chosen[0],chosen[1],r.get("SKU"),tid) in confirmed
                       for r in vrows):
                    fail("Hàng đã được KTV ký nhận tại điểm tập kết. "
                         "Không báo giao tận xã lần thứ hai.")
            rid="BC-"+uuid.uuid4().hex[:16].upper()
            submitted_at=now()
            p=lookup(read_sheet("DANH_SACH_DIEM"),"Mã điểm",chosen[1]) or {}
            dest=p.get("Link Google Maps") or ("https://www.google.com/maps/search/?api=1&query="+quote(p.get("Tên xã/phường/đơn vị",chosen[1])))
            maps=(f"https://www.google.com/maps?q={lat},{lon}" if require_gps else dest)
            st.session_state["_report_attempt_pending"]=rid
            phase="lưu ảnh Drive"
            try:
                urls=images_to_drive(
                    photos,rid,chosen[0],chosen[1],
                    p.get("Tên xã/phường/đơn vị",chosen[1]),submitted_at)
                # All source rows are refreshed after photo upload. Business
                # updates are issued together; no half-synced status/report.
                phase="đối soát dữ liệu sau tải ảnh"
                read_sheet.clear()
                vc_latest=read_sheet_live("VAN_CHUYEN")
                ld_latest=read_sheet_live("LAP_DAT")
                latest_by_tab={"VAN_CHUYEN":vc_latest,"LAP_DAT":ld_latest}
                refreshed=[]
                for tab,row in live:
                    key="Mã nhiệm vụ VC" if tab=="VAN_CHUYEN" else "Mã nhiệm vụ LD"
                    current=lookup(latest_by_tab[tab],key,row[key])
                    if (not current or
                            current.get("Trạng thái")!=row.get("Trạng thái") or
                            numeric(current.get("SL thực tế"))!=numeric(row.get("SL thực tế")) or
                            current.get("Mã đội VC" if tab=="VAN_CHUYEN" else "Mã đội KTV")!=tid):
                        raise ValueError("Nhiệm vụ thay đổi trong khi tải ảnh; tải lại và đối soát.")
                    refreshed.append(current)
                tabs=("BAO_CAO_TRIEN_KHAI","VAN_CHUYEN","LAP_DAT",
                      "TIEN_DO","NHAT_KY_HE_THONG")
                sheets={name:sheet(name).id for name in tabs}
                prior=read_sheet_live("BAO_CAO_TRIEN_KHAI")
                prior_ids={str(r.get("Mã báo cáo") or next(iter(r.values()),""))
                           for r in prior}
                requests,results=atomic_report_requests(
                    sheets,rid,submitted_at,a["username"],a["role"],tid,
                    chosen[0],chosen[1],mode,refreshed,increments,
                    urls,maps,lat,lon,vc_latest,ld_latest,
                    read_sheet_live("TIEN_DO"),
                    p.get("Tên xã/phường/đơn vị",chosen[1]),prior_ids)
                phase="ghi đồng bộ Google Sheets"
                client()[0].batch_update({"requests":requests})
                read_sheet.clear()
                st.session_state.pop("_report_attempt_pending",None)
                st.success(f"Đã ghi {len(results)} nhiệm vụ theo SL THỰC TẾ. "
                           "Báo cáo, trạng thái, tiến độ và nhật ký đã ghi cùng giao dịch.")
            except ValueError as exc:
                read_sheet.clear()
                st.error("KHÔNG ghi số liệu: "+str(exc)+
                         ". Cần đối soát trước khi gửi lại.")
            except Exception as exc:
                read_sheet.clear()
                # Google Sheets batch_update may commit despite a lost response.
                # Keep the duplicate-submit guard until ADMIN reconciles.
                error_label=(drive_error_summary(exc)
                             if phase=="lưu ảnh Drive" else type(exc).__name__)
                if phase=="ghi đồng bộ Google Sheets":
                    st.error("CHƯA XÁC NHẬN ĐƯỢC GIAO DỊCH Google Sheets ("+
                             type(exc).__name__+"). KHÔNG bấm lại. "
                             "ADMIN kiểm tra mã "+rid+
                             " trong BAO_CAO_TRIEN_KHAI rồi mới thử tiếp.")
                else:
                    st.error("CHƯA LƯU ĐƯỢC BÁO CÁO tại bước "+phase+
                             " ("+error_label+"). KHÔNG bấm lại. "
                             "ADMIN kiểm tra ảnh Drive và mã "+rid+
                             " trong BAO_CAO_TRIEN_KHAI trước khi thử tiếp.")

def dashboard():
    """Compact, read-only leadership dashboard: ADMIN / LANH_DAO."""
    require_role("ADMIN","LANH_DAO")
    st.markdown(
        """
        <style>
        [data-testid="stMainBlockContainer"] h3 {
            font-size:1.25rem !important;
            margin-top:.2rem !important;
            margin-bottom:.35rem !important;
        }
        [data-testid="stMainBlockContainer"] [data-testid="stSelectbox"] label,
        [data-testid="stMainBlockContainer"] [data-testid="stCaptionContainer"] p {
            font-size:.74rem !important;
            line-height:1.25 !important;
        }
        [data-testid="stMainBlockContainer"] [data-testid="stSelectbox"] div[data-baseweb="select"] * {
            font-size:.84rem !important;
        }
        [data-testid="stMainBlockContainer"] [data-testid="stButton"] button {
            min-height:2.15rem !important;
            padding:.3rem .55rem !important;
        }
        [data-testid="stMainBlockContainer"] [data-testid="stButton"] button p {
            font-size:.78rem !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.subheader("📊 Báo cáo tiến độ")
    projects=read_sheet("DANH_SACH_DU_AN")
    ids=[p["Mã dự án"] for p in projects if p.get("Mã dự án")]
    if not ids:
        st.info("Chưa có dự án để báo cáo.")
        return
    preferred=ids.index("DA880") if "DA880" in ids else 0
    project_col,refresh_col=st.columns([3,1])
    with project_col:
        project=st.selectbox("Dự án",ids,index=preferred,
                             key="leadership_project")
    with refresh_col:
        if st.button("🔄 CẬP NHẬT",key="leadership_refresh",
                     use_container_width=True):
            read_sheet.clear()
            st.rerun()
    progress=build_progress(
        project,projects,
        read_sheet("DANH_SACH_DIEM"),
        read_sheet("NHAP_KHO"),
        read_sheet("KHO_PHAN_BO"),
        read_sheet("VAN_CHUYEN"),
        read_sheet("LAP_DAT"),
        read_sheet("BAO_CAO_TRIEN_KHAI"),
        read_sheet("QUAN_LY_DOI"),
        read_sheet("BAN_GIAO_TAP_KET")
    )
    m=progress["metrics"]
    total=max(0,int(m.get("sites",0) or 0))
    st.caption("Số liệu lúc "+now().replace("T"," ")[:19]+
               " (giờ Việt Nam) • Nhấn CẬP NHẬT để làm mới.")
    cards=[
        ("Điểm dự án","sites"),
        ("Đã phân bổ hàng","sites_allocated"),
        ("VC giao tận xã đủ","sites_delivered"),
        ("LĐ đã lắp đủ","sites_installed"),
        ("VC bàn giao tập kết đủ","sites_vc_hub_done"),
        ("KTV nhận tập kết đủ","sites_at_hub"),
    ]
    chart_rows=[
        ("Phân bổ hàng","sites_allocated"),
        ("VC giao tận xã","sites_delivered"),
        ("LĐ lắp đặt đủ","sites_installed"),
        ("VC bàn giao tập kết","sites_vc_hub_done"),
        ("KTV nhận tập kết","sites_at_hub"),
    ]
    card_html="".join(
        '<div class="kpi"><div class="label">'+label+
        '</div><div class="value">'+str(max(0,int(m.get(key,0) or 0)))+
        '</div></div>'
        for label,key in cards
    )
    rows_html=""
    for label,key in chart_rows:
        value=max(0,int(m.get(key,0) or 0))
        pct=min(100.0,(value/total*100.0) if total else 0.0)
        rows_html+=(
            '<div class="row"><div class="row-label">'+label+
            '</div><div class="track"><div class="fill" style="width:'+
            f'{pct:.1f}%'+
            ';"></div></div><div class="row-value">'+str(value)+
            '/'+str(total)+'</div></div>'
        )
    html=(
        '<!doctype html><html><head><meta name="viewport" '
        'content="width=device-width, initial-scale=1">'
        '<style>'
        '*{box-sizing:border-box}body{margin:0;padding:2px;background:transparent;'
        'font-family:system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif;'
        'color:#1f2937}.layout{display:grid;grid-template-columns:1.05fr .95fr;'
        'gap:10px}.kpis{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));'
        'gap:8px}.kpi,.chart{border:1px solid #dbe4ee;border-radius:12px;'
        'background:#f3f7fb}.kpi{min-height:70px;padding:9px 10px;'
        'display:flex;flex-direction:column;justify-content:space-between}'
        '.label{font-size:10px;line-height:1.2;color:#475569;min-height:24px}'
        '.value{font-size:32px;line-height:1;font-weight:800;color:#111827;letter-spacing:-.02em}'
        '.chart{padding:10px}.chart-title{font-size:12px;font-weight:700;'
        'margin:0 0 7px}.row{display:grid;grid-template-columns:108px minmax(35px,1fr) 43px;'
        'gap:6px;align-items:center;margin:10px 0}.row-label{font-size:10px;'
        'line-height:1.15;color:#475569}.track{height:9px;border-radius:99px;'
        'background:#dbe5ef;overflow:hidden}.fill{height:100%;border-radius:99px;'
        'background:#2878c8}.row-value{font-size:12px;text-align:right;'
        'font-variant-numeric:tabular-nums;color:#334155}.note{font-size:9px;'
        'line-height:1.25;color:#64748b;margin-top:8px}'
        '@media(max-width:620px){.layout{grid-template-columns:1fr;gap:8px}'
        '.kpis{grid-template-columns:repeat(2,minmax(0,1fr));gap:7px}'
        '.kpi{min-height:64px;padding:7px 8px}.label{font-size:9px;min-height:22px}'
        '.value{font-size:30px}.chart{padding:8px 9px}'
        '.row{grid-template-columns:102px minmax(35px,1fr) 42px;margin:7px 0}'
        '.chart-title{margin-bottom:4px}}'
        '</style></head><body><div class="layout"><div class="kpis">'+card_html+
        '</div><div class="chart"><div class="chart-title">Tiến độ theo điểm</div>'+
        rows_html+'<div class="note">Hai số tập kết đếm xã đủ SKU theo biên bản, không phải số chuyến. '
        'Báo cáo tiến độ chưa thay nghiệm thu/đối soát.</div></div></div></body></html>'
    )
    st.components.v1.html(html,height=420,scrolling=False)

def main():
    try:
        # Public help is a separate bookmarkable page inside the PREVIEW app.
        # It has no backend reads/writes and does not bypass role permissions.
        if st.query_params.get("guide") == "1":
            render_guide()
            return
        _, guide_col = st.columns([2.5, 1.5])
        with guide_col:
            if st.button("📖 Hướng dẫn sử dụng",
                         key="open_help_from_home",
                         help="Xem cách đăng ký, giao hàng, lắp đặt và thực hành thử",
                         use_container_width=True):
                st.query_params["guide"] = "1"
                st.rerun()
        view=menu()
        if view=="__LEADERSHIP_LOGIN__":leadership_login()
        elif view=="__TEAM_LOGIN__":login(team_only=True)
        elif view=="🔐 Đăng nhập ADMIN / tài khoản khác":login()
        elif view=="📝 Đăng ký tham gia":register()
        elif view=="📝 Đăng ký hộ đội trưởng":register(on_behalf=True)
        elif view=="🗺️ Tìm đường":directions()
        elif view=="📁 Dự án/Điểm/Nhập kho":projects_and_inventory()
        elif view=="🧹 RESET / XÓA DỰ ÁN":project_cleanup_admin(
            require_role, valid_login, client, read_sheet, audit)
        elif view=="📦 Phân bổ":allocate()
        elif view=="👥 Đăng ký & quản lý đội":teams_and_approval()
        elif view=="📦 Bàn giao tập kết":hub_admin()
        elif view=="📋 Báo cáo nhiệm vụ":report()
        elif view=="📊 Báo cáo tiến độ":dashboard()
    except Exception as e:
        if isinstance(e, PreviewSheetReadError):
            st.error(str(e))
        elif isinstance(e, gspread.exceptions.APIError):
            st.error("Không thể kết nối Google Sheets DỰ ÁN: HTTP " +
                     str(_google_api_status(e) or "không rõ"))
        else:
            st.error("Không thể truy cập dữ liệu DỰ ÁN: " + type(e).__name__)
        st.caption("Không ghi hoặc xóa dữ liệu khi bước đọc thất bại; kiểm tra mã HTTP nếu còn lỗi.")
if __name__=="__main__":
    main()
