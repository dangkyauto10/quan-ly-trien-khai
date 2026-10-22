"""Safe XLSX reader for multi-SKU location exceptions (no Sheet writes)."""
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from preview_bulk_allocation import quantity


def parse_exception_excel(filename, data, allowed_ids, skus):
    """Return (point_sku_overrides, affected_point_ids).

    Headers must have Mã điểm and one or more exact SKU columns. A blank cell
    keeps the default amount; explicit 0 disables allocation for that SKU.
    """
    if Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("Chỉ chấp nhận file Excel .xlsx.")
    if not data or len(data) > 5_000_000:
        raise ValueError("File trống hoặc vượt quá 5 MB.")
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        head = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not head:
            raise ValueError("File Excel chưa có dòng tiêu đề.")
        head = [str(c or "").strip() for c in head[:100]]
        if "Mã điểm" not in head:
            raise ValueError("Excel phải có cột 'Mã điểm'.")
        point_col = head.index("Mã điểm")
        sku_cols = {j:h for j,h in enumerate(head) if h in skus}
        if not sku_cols:
            raise ValueError("Excel cần ít nhất một cột có tên SKU của dự án.")
        allowed = set(allowed_ids)
        seen, overr, points = set(), {}, []
        for row in ws.iter_rows(min_row=2, max_row=2001, max_col=len(head), values_only=True):
            pid = str(row[point_col] or "").strip()
            if not pid:
                continue
            if pid not in allowed:
                raise ValueError(f"Mã điểm {pid} không thuộc danh sách đang chọn.")
            if pid in seen:
                raise ValueError(f"Mã điểm bị trùng trong Excel: {pid}.")
            seen.add(pid)
            points.append(pid)
            for idx, sku in sku_cols.items():
                value = row[idx] if idx < len(row) else None
                if value is None or str(value).strip() == "":
                    continue
                overr[(pid,sku)] = float(quantity(value))
        if not points:
            raise ValueError("Không có địa điểm hợp lệ trong file.")
        return overr, points
    finally:
        wb.close()
