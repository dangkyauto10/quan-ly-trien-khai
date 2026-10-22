"""Read candidate project locations from Excel .xlsx or Word .docx.

Preview-only: parsing never writes to Google Sheets; the Admin must review
candidates and explicitly confirm the append operation.
"""
from io import BytesIO
import re
import unicodedata
from pathlib import Path

from docx import Document
from openpyxl import load_workbook, Workbook


MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS_PER_SHEET = 2000
MAX_COLS = 24

_HEADER_TOKENS = (
    "tên địa điểm", "địa điểm", "địa điểm giao hàng và lắp đặt",
    "địa điểm giao hàng", "tên xã/phường/đơn vị", "tên xã phường đơn vị",
    "xã/phường/đơn vị", "tên xã", "đơn vị triển khai",
    "tên đơn vị", "danh sách địa điểm"
)


def normalize_name(value):
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).casefold().split())


def is_header(text):
    return normalize_name(text).strip(":") in _HEADER_TOKENS


def column_choice(rows):
    """Choose the first identified location header; otherwise the first column."""
    for row in rows[:5]:
        for i, cell in enumerate(row):
            if is_header(cell):
                return i
    return 0


def read_location_file(filename, content):
    if not isinstance(content, bytes) or not content:
        raise ValueError("Tệp tải lên trống.")
    if len(content) > MAX_BYTES:
        raise ValueError("Tệp vượt quá 5 MB. Hãy tách nhỏ danh sách.")
    ext = Path(filename).suffix.lower()
    sources = {}
    if ext == ".xlsx":
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            for sh in wb.worksheets:
                rows = []
                for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, MAX_ROWS_PER_SHEET),
                                        max_col=min(sh.max_column, MAX_COLS), values_only=True):
                    rows.append([str(c).strip() if c is not None else "" for c in row])
                sources[sh.title] = rows
        finally:
            wb.close()
    elif ext == ".docx":
        doc = Document(BytesIO(content))
        rows = []
        for paragraph in doc.paragraphs:
            for line in paragraph.text.splitlines():
                value = line.strip()
                if value:
                    rows.append([value])
        for table in doc.tables:
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells[:MAX_COLS]])
        sources["Tài liệu Word"] = rows[:MAX_ROWS_PER_SHEET]
    else:
        raise ValueError("Chỉ hỗ trợ Word .docx và Excel .xlsx.")
    return {name: rows for name, rows in sources.items() if rows}


def candidate_names(rows, column_index=0, skip_first=False):
    result = []
    for row_index, row in enumerate(rows):
        if row_index == 0 and skip_first:
            continue
        if column_index >= len(row):
            continue
        name = re.sub(r"^\s*\d{1,4}[\.\)]\s+", "", str(row[column_index])).strip()
        name = " ".join(name.split())
        if not name or is_header(name):
            continue
        if len(name) > 180:
            continue
        result.append(name)
    return result


def blank_excel_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "DANH_SACH_DIEM"
    ws["A1"] = "Tên địa điểm"
    ws.column_dimensions["A"].width = 44
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
