import io
import unittest

from docx import Document
from openpyxl import Workbook

from preview_import_locations import (
    read_location_file, candidate_names, column_choice, normalize_name,
    blank_excel_template,
)


class PreviewLocationImportTest(unittest.TestCase):
    def test_excel_single_column(self):
        wb=Workbook()
        ws=wb.active
        ws.append(["Tên địa điểm"])
        ws.append(["Phường Minh Xuân"])
        ws.append(["Xã Yên Sơn"])
        b=io.BytesIO()
        wb.save(b)
        tabs=read_location_file("diem.xlsx", b.getvalue())
        rows=tabs[ws.title]
        self.assertEqual(column_choice(rows),0)
        self.assertEqual(candidate_names(rows,0,True),["Phường Minh Xuân","Xã Yên Sơn"])

    def test_excel_existing_sheet_column(self):
        wb=Workbook()
        ws=wb.active
        ws.append(["Tuyến","Thứ tự","Địa điểm giao hàng và lắp đặt"])
        ws.append(["T01",1,"Phường Minh Xuân"])
        b=io.BytesIO()
        wb.save(b)
        rows=list(read_location_file("diem.xlsx", b.getvalue()).values())[0]
        col=column_choice(rows)
        self.assertEqual(col,2)
        self.assertEqual(candidate_names(rows,col),["Phường Minh Xuân"])

    def test_word_paragraph_list(self):
        doc=Document()
        doc.add_paragraph("1. Phường Hà Giang 1")
        doc.add_paragraph("2. Phường Hà Giang 2")
        b=io.BytesIO()
        doc.save(b)
        rows=read_location_file("danh-sach.docx",b.getvalue())["Tài liệu Word"]
        self.assertEqual(candidate_names(rows),["Phường Hà Giang 1","Phường Hà Giang 2"])

    def test_word_table(self):
        doc=Document()
        table=doc.add_table(rows=1,cols=2)
        table.rows[0].cells[0].text="STT"
        table.rows[0].cells[1].text="Tên địa điểm"
        row=table.add_row()
        row.cells[0].text="1"
        row.cells[1].text="Xã Sơn Dương"
        b=io.BytesIO()
        doc.save(b)
        rows=read_location_file("data.docx",b.getvalue())["Tài liệu Word"]
        col=column_choice(rows)
        self.assertEqual(col,1)
        self.assertEqual(candidate_names(rows,col),["Xã Sơn Dương"])

    def test_blank_template_and_unicode(self):
        rows=read_location_file("mau.xlsx",blank_excel_template())
        self.assertTrue(rows)
        self.assertEqual(normalize_name("  PHƯỜNG   MINH Xuân "),normalize_name("Phường Minh Xuân"))


if __name__=="__main__":
    unittest.main()
